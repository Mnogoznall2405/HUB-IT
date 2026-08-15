"""Benchmark RouterAI models on large CSV reconciliation and XLSX creation.

The models receive the same deterministic Russian-language multi-file dataset
and return a constrained workbook specification. The harness builds a real
Excel workbook from that specification and validates both the analysis and the
resulting file without executing model-generated code.
"""
from __future__ import annotations

import argparse
import csv
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta
import io
import json
import math
from pathlib import Path
import re
import shutil
import subprocess
import sys
import time
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.benchmark_routerai_corporate import _extract_completion_text, _parse_json  # noqa: E402
from shared.llm.client import OpenRouterClient  # noqa: E402


DEPARTMENTS = (
    "ИТ",
    "Финансы",
    "Закупки",
    "Продажи",
    "Логистика",
    "Производство",
    "Юристы",
    "HR",
)
LOCATIONS = tuple(f"Площадка-{index:02d}" for index in range(1, 11))
CATEGORIES = ("Ноутбук", "Монитор", "Принтер", "Сервер", "Телефон")
REQUIRED_SHEETS = ("Сводка", "Отклонения", "Активы", "Перемещения", "Сотрудники")


def _csv_text(headers: list[str], rows: list[dict[str, Any]]) -> str:
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=headers, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue()


def _build_dataset() -> dict[str, Any]:
    employees = []
    for index in range(1, 801):
        employees.append(
            {
                "employee_id": f"E{index:04d}",
                "full_name": f"Сотрудник {index:04d}",
                "department": DEPARTMENTS[(index - 1) % len(DEPARTMENTS)],
                "status": "inactive" if 791 <= index <= 800 else "active",
            }
        )

    duplicate_numbers = set(range(101, 116))
    inactive_numbers = set(range(1001, 1021))
    missing_numbers = set(range(2001, 2013))
    mismatch_numbers = set(range(3001, 3031))

    assets = []
    canonical_assets: dict[str, dict[str, Any]] = {}
    latest_locations: dict[str, str] = {}
    for index in range(1, 5001):
        asset_id = f"A{index:06d}"
        latest_location = LOCATIONS[(index * 3) % len(LOCATIONS)]
        latest_locations[asset_id] = latest_location
        current_location = (
            LOCATIONS[(index * 3 + 1) % len(LOCATIONS)]
            if index in mismatch_numbers
            else latest_location
        )
        assigned_to = f"E{((index - 1) % 780) + 1:04d}"
        if index in inactive_numbers:
            assigned_to = f"E{791 + ((index - 1001) % 10):04d}"
        elif index in missing_numbers:
            assigned_to = f"E{9001 + (index - 2001):04d}"
        row = {
            "asset_id": asset_id,
            "serial_number": f"SN{index:08d}",
            "category": CATEGORIES[(index - 1) % len(CATEGORIES)],
            "department": DEPARTMENTS[(index - 1) % len(DEPARTMENTS)],
            "assigned_to": assigned_to,
            "current_location": current_location,
            "purchase_date": (date(2019, 1, 1) + timedelta(days=(index * 17) % 2555)).isoformat(),
            "book_value_rub": round(15000 + (index * 137) % 285000, 2),
        }
        assets.append(row)
        canonical_assets[asset_id] = row

    for index in sorted(duplicate_numbers):
        duplicate = dict(canonical_assets[f"A{index:06d}"])
        duplicate["serial_number"] += "-COPY"
        assets.append(duplicate)

    movements = []
    movement_number = 1
    for index in range(1, 5001):
        asset_id = f"A{index:06d}"
        destination = latest_locations[asset_id]
        movements.append(
            {
                "movement_id": f"M{movement_number:07d}",
                "asset_id": asset_id,
                "moved_at": f"2026-07-{(index % 28) + 1:02d}T12:00:00",
                "from_location": LOCATIONS[(index * 3 + 7) % len(LOCATIONS)],
                "to_location": destination,
                "approved_by": f"E{((index + 40) % 780) + 1:04d}",
            }
        )
        movement_number += 1
    for index in range(1, 981):
        movements.append(
            {
                "movement_id": f"M{movement_number:07d}",
                "asset_id": f"A{index:06d}",
                "moved_at": f"2026-06-{(index % 28) + 1:02d}T09:00:00",
                "from_location": LOCATIONS[(index + 2) % len(LOCATIONS)],
                "to_location": LOCATIONS[(index + 3) % len(LOCATIONS)],
                "approved_by": f"E{((index + 80) % 780) + 1:04d}",
            }
        )
        movement_number += 1
    orphan_ids = [f"A{900000 + index:06d}" for index in range(1, 21)]
    for index, asset_id in enumerate(orphan_ids, 1):
        movements.append(
            {
                "movement_id": f"M{movement_number:07d}",
                "asset_id": asset_id,
                "moved_at": f"2026-07-{index:02d}T15:00:00",
                "from_location": "Площадка-01",
                "to_location": "Площадка-02",
                "approved_by": f"E{index:04d}",
            }
        )
        movement_number += 1

    department_summary = []
    for department in DEPARTMENTS:
        rows = [row for row in canonical_assets.values() if row["department"] == department]
        department_summary.append(
            {
                "department": department,
                "asset_count": len(rows),
                "book_value_rub": round(sum(float(row["book_value_rub"]) for row in rows), 2),
            }
        )

    expected = {
        "raw_asset_rows": len(assets),
        "unique_asset_ids": len(canonical_assets),
        "movement_rows": len(movements),
        "employee_rows": len(employees),
        "duplicate_asset_ids": [f"A{index:06d}" for index in sorted(duplicate_numbers)],
        "orphan_movement_asset_ids": orphan_ids,
        "inactive_employee_asset_ids": [f"A{index:06d}" for index in sorted(inactive_numbers)],
        "missing_employee_asset_ids": [f"A{index:06d}" for index in sorted(missing_numbers)],
        "location_mismatch_asset_ids": [f"A{index:06d}" for index in sorted(mismatch_numbers)],
        "total_book_value_rub": round(
            sum(float(row["book_value_rub"]) for row in canonical_assets.values()), 2
        ),
        "department_summary": department_summary,
    }
    expected["exception_count"] = sum(
        len(expected[key])
        for key in (
            "duplicate_asset_ids",
            "orphan_movement_asset_ids",
            "inactive_employee_asset_ids",
            "missing_employee_asset_ids",
            "location_mismatch_asset_ids",
        )
    )

    files = {
        "assets.csv": _csv_text(list(assets[0]), assets),
        "movements.csv": _csv_text(list(movements[0]), movements),
        "employees.csv": _csv_text(list(employees[0]), employees),
    }
    return {
        "assets": assets,
        "movements": movements,
        "employees": employees,
        "files": files,
        "expected": expected,
    }


def _build_prompt(files: dict[str, str]) -> str:
    file_blocks = "\n\n".join(
        f'<file name="{name}" format="csv">\n{content}</file>' for name, content in files.items()
    )
    return f"""Ты анализируешь три взаимосвязанных корпоративных CSV-файла и проектируешь итоговую Excel-книгу. Используй только данные файлов. Не пересказывай исходные строки и не добавляй факты.

Правила сверки:
1. Канонической записью актива считается первое появление asset_id в assets.csv. Все последующие появления того же ID — дубли.
2. Сиротское перемещение — asset_id из movements.csv, которого нет среди канонических активов.
3. Неактивное назначение — канонический актив назначен сотруднику со status=inactive.
4. Отсутствующий сотрудник — assigned_to канонического актива отсутствует в employees.csv.
5. Расхождение площадки — current_location канонического актива не совпадает с to_location его самого позднего перемещения по moved_at. Актив без перемещений сюда не относится.
6. Сводка по подразделениям и балансовая стоимость считаются только по каноническим активам.

Требования к Excel:
- листы строго в порядке: Сводка, Отклонения, Активы, Перемещения, Сотрудники;
- на листе Активы в колонке I должна быть копируемая формула-флаг Канон/Дубликат;
- Сводка должна использовать формулы Excel для количества строк, уникальных активов, дублей, общей стоимости и числа отклонений;
- таблица подразделений начинается с A11: подразделение, количество активов, балансовая стоимость;
- формулы должны быть совместимы с обычным Excel: английские имена функций, разделитель аргументов запятая, без FILTER/XLOOKUP/SORT/SEQUENCE/INDIRECT/OFFSET;
- нужны freeze panes, автофильтры на табличных листах, денежный и датовый форматы, условное выделение отклонений и столбчатая диаграмма стоимости по подразделениям.

Верни только один валидный JSON без markdown и пояснений по схеме:
{{
  "analysis": {{
    "raw_asset_rows": number,
    "unique_asset_ids": number,
    "movement_rows": number,
    "employee_rows": number,
    "duplicate_asset_ids": [string],
    "orphan_movement_asset_ids": [string],
    "inactive_employee_asset_ids": [string],
    "missing_employee_asset_ids": [string],
    "location_mismatch_asset_ids": [string],
    "total_book_value_rub": number,
    "department_summary": [{{"department": string, "asset_count": number, "book_value_rub": number}}]
  }},
  "excel_plan": {{
    "sheet_order": [string],
    "freeze_panes": {{"Сводка": string, "Отклонения": string, "Активы": string, "Перемещения": string, "Сотрудники": string}},
    "autofilter_sheets": [string],
    "asset_record_flag_formula": string,
    "summary_formulas": {{
      "raw_asset_rows": string,
      "unique_asset_ids": string,
      "duplicate_rows": string,
      "total_book_value_rub": string,
      "exception_count": string
    }},
    "department_count_formula": string,
    "department_value_formula": string,
    "header_fill": string,
    "header_font_color": string,
    "currency_format": string,
    "date_format": string,
    "conditional_format": boolean,
    "chart": {{"type": "bar", "title": string}}
  }}
}}

Формулы для подразделений дай для первой строки данных 11; при копировании вниз ссылки должны изменяться корректно.

{file_blocks}"""


def _safe_formula(value: Any) -> str | None:
    if not isinstance(value, str) or not value.startswith("=") or len(value) > 1000:
        return None
    upper = value.upper()
    if any(token in upper for token in ("WEBSERVICE", "HYPERLINK", "INDIRECT", "OFFSET", "[", "]")):
        return None
    return value


def _hex_color(value: Any, fallback: str) -> str:
    text = str(value or "").strip().lstrip("#").upper()
    return text if re.fullmatch(r"[0-9A-F]{6}", text) else fallback


def _normalize_ids(value: Any) -> set[str]:
    if not isinstance(value, list):
        return set()
    return {str(item).strip().upper() for item in value}


def _number_equal(value: Any, expected: float) -> bool:
    try:
        return math.isclose(float(value), float(expected), rel_tol=0, abs_tol=0.01)
    except (TypeError, ValueError):
        return False


def _analysis_score(parsed: dict[str, Any] | None, expected: dict[str, Any]) -> tuple[int, dict[str, int]]:
    analysis = parsed.get("analysis") if isinstance(parsed, dict) else None
    if not isinstance(analysis, dict):
        return 0, {"counts": 0, "sets": 0, "totals": 0}

    counts = sum(
        4 if _number_equal(analysis.get(key), expected[key]) else 0
        for key in ("raw_asset_rows", "unique_asset_ids", "movement_rows", "employee_rows")
    )
    set_weights = {
        "duplicate_asset_ids": 12,
        "orphan_movement_asset_ids": 12,
        "inactive_employee_asset_ids": 12,
        "missing_employee_asset_ids": 8,
        "location_mismatch_asset_ids": 15,
    }
    sets = sum(
        weight
        if _normalize_ids(analysis.get(key)) == _normalize_ids(expected[key])
        else 0
        for key, weight in set_weights.items()
    )
    totals = 5 if _number_equal(analysis.get("total_book_value_rub"), expected["total_book_value_rub"]) else 0

    actual_departments = {
        str(row.get("department")): row
        for row in analysis.get("department_summary", [])
        if isinstance(row, dict)
    }
    expected_departments = {row["department"]: row for row in expected["department_summary"]}
    if set(actual_departments) == set(expected_departments):
        if all(
            _number_equal(actual_departments[name].get("asset_count"), expected_departments[name]["asset_count"])
            for name in expected_departments
        ):
            totals += 2
        if all(
            _number_equal(actual_departments[name].get("book_value_rub"), expected_departments[name]["book_value_rub"])
            for name in expected_departments
        ):
            totals += 3
    return counts + sets + totals, {"counts": counts, "sets": sets, "totals": totals}


def _write_table(sheet: Any, headers: list[str], rows: list[dict[str, Any]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])


def _apply_header_style(sheet: Any, fill_color: str, font_color: str, row: int = 1) -> None:
    for cell in sheet[row]:
        if cell.value is not None:
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.font = Font(color=font_color, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")


def _build_workbook(
    parsed: dict[str, Any] | None,
    dataset: dict[str, Any],
    output_path: Path,
) -> dict[str, Any]:
    parsed = parsed if isinstance(parsed, dict) else {}
    analysis = parsed.get("analysis") if isinstance(parsed.get("analysis"), dict) else {}
    plan = parsed.get("excel_plan") if isinstance(parsed.get("excel_plan"), dict) else {}
    fill = _hex_color(plan.get("header_fill"), "1F4E78")
    font_color = _hex_color(plan.get("header_font_color"), "FFFFFF")
    currency_format = str(plan.get("currency_format") or '#,##0.00 [$₽-ru-RU]')
    date_format = str(plan.get("date_format") or "dd.mm.yyyy")

    workbook = Workbook()
    workbook.remove(workbook.active)
    requested_order = plan.get("sheet_order") if isinstance(plan.get("sheet_order"), list) else []
    order = [name for name in requested_order if name in REQUIRED_SHEETS]
    for name in REQUIRED_SHEETS:
        if name not in order:
            order.append(name)
    sheets = {name: workbook.create_sheet(name) for name in order}

    assets_sheet = sheets["Активы"]
    asset_headers = list(dataset["assets"][0])
    _write_table(assets_sheet, asset_headers, dataset["assets"])
    assets_sheet.cell(1, 9, "record_flag")
    flag_formula = _safe_formula(plan.get("asset_record_flag_formula"))
    if flag_formula:
        assets_sheet.cell(2, 9, flag_formula)
        for row_number in range(3, assets_sheet.max_row + 1):
            try:
                translated = Translator(flag_formula, origin="I2").translate_formula(f"I{row_number}")
            except Exception:  # noqa: BLE001 - malformed model formula remains observable
                translated = flag_formula
            assets_sheet.cell(row_number, 9, translated)

    movements_sheet = sheets["Перемещения"]
    _write_table(movements_sheet, list(dataset["movements"][0]), dataset["movements"])
    employees_sheet = sheets["Сотрудники"]
    _write_table(employees_sheet, list(dataset["employees"][0]), dataset["employees"])

    exceptions_sheet = sheets["Отклонения"]
    exceptions_sheet.append(["type", "asset_id", "detail"])
    exception_keys = (
        ("duplicate_asset_ids", "duplicate_asset_id"),
        ("orphan_movement_asset_ids", "orphan_movement"),
        ("inactive_employee_asset_ids", "inactive_employee"),
        ("missing_employee_asset_ids", "missing_employee"),
        ("location_mismatch_asset_ids", "location_mismatch"),
    )
    for key, exception_type in exception_keys:
        values = analysis.get(key) if isinstance(analysis.get(key), list) else []
        for asset_id in values:
            exceptions_sheet.append([exception_type, str(asset_id), key])

    summary_sheet = sheets["Сводка"]
    summary_sheet.merge_cells("A1:C1")
    summary_sheet["A1"] = "Сверка корпоративных активов"
    summary_sheet["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    summary_sheet["A1"].fill = PatternFill("solid", fgColor=fill)
    summary_sheet["A1"].alignment = Alignment(horizontal="center")
    metrics = (
        ("Строк активов", "raw_asset_rows"),
        ("Уникальных активов", "unique_asset_ids"),
        ("Дублирующих строк", "duplicate_rows"),
        ("Балансовая стоимость, ₽", "total_book_value_rub"),
        ("Отклонений", "exception_count"),
    )
    formulas = plan.get("summary_formulas") if isinstance(plan.get("summary_formulas"), dict) else {}
    for row_number, (label, key) in enumerate(metrics, 3):
        summary_sheet.cell(row_number, 1, label)
        summary_sheet.cell(row_number, 2, _safe_formula(formulas.get(key)))
    summary_sheet["B6"].number_format = currency_format

    summary_sheet["A9"] = "Подразделения"
    summary_sheet["A9"].font = Font(bold=True)
    summary_sheet.append([])
    summary_sheet["A10"] = "Подразделение"
    summary_sheet["B10"] = "Количество активов"
    summary_sheet["C10"] = "Балансовая стоимость, ₽"
    departments = analysis.get("department_summary") if isinstance(analysis.get("department_summary"), list) else []
    for index, row in enumerate(departments, 11):
        if not isinstance(row, dict):
            continue
        summary_sheet.cell(index, 1, row.get("department"))
        count_formula = _safe_formula(plan.get("department_count_formula"))
        value_formula = _safe_formula(plan.get("department_value_formula"))
        if count_formula:
            try:
                count_formula = Translator(count_formula, origin="B11").translate_formula(f"B{index}")
            except Exception:  # noqa: BLE001
                pass
        if value_formula:
            try:
                value_formula = Translator(value_formula, origin="C11").translate_formula(f"C{index}")
            except Exception:  # noqa: BLE001
                pass
        summary_sheet.cell(index, 2, count_formula)
        summary_sheet.cell(index, 3, value_formula).number_format = currency_format

    for sheet in workbook.worksheets:
        _apply_header_style(sheet, fill, font_color, 1)
        sheet.sheet_view.showGridLines = False
    _apply_header_style(summary_sheet, fill, font_color, 10)

    freeze_panes = plan.get("freeze_panes") if isinstance(plan.get("freeze_panes"), dict) else {}
    for name, cell in freeze_panes.items():
        if name in sheets and isinstance(cell, str) and re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", cell):
            sheets[name].freeze_panes = cell

    autofilters = set(plan.get("autofilter_sheets") or [])
    for name in autofilters:
        if name in sheets and name != "Сводка":
            sheets[name].auto_filter.ref = sheets[name].dimensions

    for row in range(2, assets_sheet.max_row + 1):
        if isinstance(assets_sheet.cell(row, 7).value, str):
            assets_sheet.cell(row, 7).value = date.fromisoformat(assets_sheet.cell(row, 7).value)
        assets_sheet.cell(row, 7).number_format = date_format
        assets_sheet.cell(row, 8).number_format = currency_format
    for row in range(2, movements_sheet.max_row + 1):
        if isinstance(movements_sheet.cell(row, 3).value, str):
            movements_sheet.cell(row, 3).value = datetime.fromisoformat(movements_sheet.cell(row, 3).value)
        movements_sheet.cell(row, 3).number_format = "dd.mm.yyyy hh:mm"

    if plan.get("conditional_format") is True and exceptions_sheet.max_row >= 2:
        exceptions_sheet.conditional_formatting.add(
            f"A2:C{exceptions_sheet.max_row}",
            FormulaRule(formula=['$A2="location_mismatch"'], fill=PatternFill("solid", fgColor="FCE4D6")),
        )

    chart = plan.get("chart") if isinstance(plan.get("chart"), dict) else {}
    if chart.get("type") == "bar" and summary_sheet.max_row >= 11:
        bar_chart = BarChart()
        bar_chart.type = "bar"
        bar_chart.title = str(chart.get("title") or "Балансовая стоимость по подразделениям")
        bar_chart.y_axis.title = "Подразделение"
        bar_chart.x_axis.title = "₽"
        data = Reference(summary_sheet, min_col=3, min_row=10, max_row=summary_sheet.max_row)
        categories = Reference(summary_sheet, min_col=1, min_row=11, max_row=summary_sheet.max_row)
        bar_chart.add_data(data, titles_from_data=True)
        bar_chart.set_categories(categories)
        bar_chart.height = 8
        bar_chart.width = 14
        summary_sheet.add_chart(bar_chart, "E3")

    widths = {
        "Сводка": {"A": 30, "B": 20, "C": 24},
        "Отклонения": {"A": 25, "B": 16, "C": 34},
        "Активы": {"A": 14, "B": 18, "C": 16, "D": 18, "E": 14, "F": 16, "G": 14, "H": 18, "I": 14},
        "Перемещения": {"A": 14, "B": 14, "C": 20, "D": 16, "E": 16, "F": 14},
        "Сотрудники": {"A": 14, "B": 24, "C": 18, "D": 12},
    }
    for sheet_name, sheet_widths in widths.items():
        for column, width in sheet_widths.items():
            sheets[sheet_name].column_dimensions[column].width = width

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    return _validate_workbook(output_path, plan)


def _formula_contains(formula: Any, *tokens: str) -> bool:
    if not isinstance(formula, str) or not formula.startswith("="):
        return False
    upper = re.sub(r"\s+", "", formula.upper())
    return all(token.upper() in upper for token in tokens)


def _validate_workbook(path: Path, plan: dict[str, Any]) -> dict[str, Any]:
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:  # noqa: BLE001
        return {"opens": False, "error": f"{type(exc).__name__}: {exc}", "score": 0}

    sheet_order = workbook.sheetnames
    summary = workbook["Сводка"] if "Сводка" in workbook else None
    assets = workbook["Активы"] if "Активы" in workbook else None
    checks = {
        "opens": True,
        "sheet_order": sheet_order == list(REQUIRED_SHEETS)
        and plan.get("sheet_order") == list(REQUIRED_SHEETS),
        "record_flag_formula": bool(
            assets and _formula_contains(assets["I2"].value, "COUNTIF", "КАНОН", "ДУБЛИКАТ")
        ),
        "summary_formulas": bool(
            summary and all(isinstance(summary[cell].value, str) and summary[cell].value.startswith("=") for cell in ("B3", "B4", "B5", "B6", "B7"))
        ),
        "department_formulas": bool(
            summary and _formula_contains(summary["B11"].value, "COUNTIFS") and _formula_contains(summary["C11"].value, "SUMIFS")
        ),
        "freeze_panes": all(workbook[name].freeze_panes for name in REQUIRED_SHEETS),
        "autofilters": all(workbook[name].auto_filter.ref for name in REQUIRED_SHEETS[1:]),
        "conditional_format": bool(workbook["Отклонения"].conditional_formatting),
        "chart": bool(summary and summary._charts),
        "formatting_spec": all(
            isinstance(plan.get(key), str) and bool(str(plan.get(key)).strip())
            for key in ("header_fill", "header_font_color", "currency_format", "date_format")
        ),
    }
    score = 0
    score += 1 if checks["opens"] else 0
    score += 2 if checks["sheet_order"] else 0
    score += 2 if checks["record_flag_formula"] else 0
    score += 3 if checks["summary_formulas"] else 0
    score += 2 if checks["department_formulas"] else 0
    score += 1 if checks["freeze_panes"] else 0
    score += 1 if checks["autofilters"] else 0
    score += 1 if checks["conditional_format"] else 0
    score += 1 if checks["chart"] else 0
    score += 1 if checks["formatting_spec"] else 0
    checks["score"] = score
    checks["max_score"] = 15
    workbook.close()
    return checks


def _recalculate_and_render(path: Path, preview_dir: Path) -> dict[str, Any]:
    soffice_candidates = (
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
    )
    soffice = next((candidate for candidate in soffice_candidates if candidate.exists()), None)
    if soffice is None:
        return {"available": False, "recalculated": False, "pdf": None, "png": None}

    preview_dir.mkdir(parents=True, exist_ok=True)
    creation_flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    pdf_result = subprocess.run(
        [str(soffice), "--headless", "--convert-to", "pdf", "--outdir", str(preview_dir), str(path)],
        capture_output=True,
        text=True,
        timeout=180,
        creationflags=creation_flags,
        check=False,
    )
    pdf_path = preview_dir / f"{path.stem}.pdf"
    png_path = None
    if pdf_path.exists():
        try:
            import fitz  # type: ignore

            document = fitz.open(pdf_path)
            page = document.load_page(0)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            png_path = preview_dir / f"{path.stem}_page1.png"
            pixmap.save(png_path)
            document.close()
        except Exception:  # noqa: BLE001 - PDF remains available for manual review
            png_path = None
    return {
        "available": True,
        "recalculated": pdf_result.returncode == 0,
        "pdf": str(pdf_path) if pdf_path.exists() else None,
        "png": str(png_path) if png_path and png_path.exists() else None,
        "stdout": pdf_result.stdout[-1000:],
        "stderr": pdf_result.stderr[-1000:],
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--models", nargs="+", required=True)
    parser.add_argument("--output-dir", type=Path, default=PROJECT_ROOT / "output" / "spreadsheet")
    parser.add_argument("--temp-dir", type=Path, default=PROJECT_ROOT / "tmp" / "spreadsheets" / "routerai_excel_benchmark")
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--timeout", type=float, default=900.0)
    parser.add_argument("--max-tokens", type=int, default=32000)
    parser.add_argument("--seed", type=int, default=20260815)
    parser.add_argument("--reasoning-effort")
    args = parser.parse_args()

    dataset = _build_dataset()
    args.temp_dir.mkdir(parents=True, exist_ok=True)
    for name, content in dataset["files"].items():
        (args.temp_dir / name).write_text(content, encoding="utf-8", newline="")
    prompt = _build_prompt(dataset["files"])

    base_client = OpenRouterClient()._build_client(timeout=args.timeout)
    client = base_client.with_options(timeout=args.timeout, max_retries=0)
    catalog = {item.id: item.model_dump() for item in client.models.list().data}
    missing = [model for model in args.models if model not in catalog]
    if missing:
        parser.error(f"models not present in RouterAI catalog: {', '.join(missing)}")

    def run_model(model: str) -> dict[str, Any]:
        supported = set(catalog[model].get("supported_parameters") or [])
        kwargs: dict[str, Any] = {
            "model": model,
            "messages": [
                {
                    "role": "system",
                    "content": "Ты выполняешь проверяемую межфайловую сверку и проектируешь Excel. Верни только итоговый JSON.",
                },
                {"role": "user", "content": prompt},
            ],
            "max_tokens": args.max_tokens,
        }
        if "temperature" in supported:
            kwargs["temperature"] = 0
        if "seed" in supported:
            kwargs["seed"] = args.seed
        if "response_format" in supported:
            kwargs["response_format"] = {"type": "json_object"}
        if args.reasoning_effort:
            kwargs["reasoning_effort"] = args.reasoning_effort

        started = time.perf_counter()
        try:
            response = client.chat.completions.create(**kwargs)
            raw = _extract_completion_text(response)
            usage = response.usage.model_dump() if getattr(response, "usage", None) else None
            parsed = _parse_json(raw)
            error = None if parsed is not None and usage is not None else "invalid_json_or_missing_usage"
        except Exception as exc:  # noqa: BLE001 - provider failures are benchmark output
            raw = ""
            usage = None
            parsed = None
            error = f"{type(exc).__name__}: {exc}"
        seconds = round(time.perf_counter() - started, 3)

        slug = re.sub(r"[^a-z0-9]+", "_", model.lower()).strip("_")
        workbook_path = args.output_dir / f"routerai_excel_{slug}.xlsx"
        workbook_checks = _build_workbook(parsed, dataset, workbook_path)
        preview = _recalculate_and_render(workbook_path, args.output_dir / "previews")
        analysis_score, score_parts = _analysis_score(parsed, dataset["expected"])
        workbook_score = int(workbook_checks.get("score") or 0)
        total_score = analysis_score + workbook_score
        return {
            "model": model,
            "reasoning_effort": args.reasoning_effort,
            "ok": parsed is not None and usage is not None,
            "score": total_score,
            "max_score": 100,
            "analysis_score": analysis_score,
            "analysis_max_score": 85,
            "workbook_score": workbook_score,
            "workbook_max_score": 15,
            "score_parts": score_parts,
            "seconds": seconds,
            "usage": usage,
            "error": error,
            "parsed": parsed,
            "raw": raw,
            "workbook": str(workbook_path),
            "workbook_checks": workbook_checks,
            "preview": preview,
        }

    print(
        f"START models={len(args.models)} prompt_chars={len(prompt)} "
        f"asset_rows={len(dataset['assets'])} movement_rows={len(dataset['movements'])} employee_rows={len(dataset['employees'])}",
        flush=True,
    )
    results = []
    checkpoint = args.report.with_suffix(".jsonl")
    checkpoint.parent.mkdir(parents=True, exist_ok=True)
    checkpoint.write_text("", encoding="utf-8")
    with ThreadPoolExecutor(max_workers=len(args.models)) as pool:
        futures = {pool.submit(run_model, model): model for model in args.models}
        for future in as_completed(futures):
            result = future.result()
            results.append(result)
            with checkpoint.open("a", encoding="utf-8") as handle:
                handle.write(json.dumps(result, ensure_ascii=False) + "\n")
            print(
                f"RESULT {result['model']} score={result['score']}/{result['max_score']} "
                f"seconds={result['seconds']} cost={(result.get('usage') or {}).get('cost')}",
                flush=True,
            )

    results.sort(key=lambda row: (-row["score"], row["seconds"]))
    args.report.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "generated_at": datetime.now().astimezone().isoformat(),
        "method": "large_multifile_csv_to_xlsx_ru",
        "reasoning_effort": args.reasoning_effort,
        "dataset": {
            "prompt_chars": len(prompt),
            "file_chars": {name: len(content) for name, content in dataset["files"].items()},
            "expected": dataset["expected"],
        },
        "results": results,
    }
    args.report.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"REPORT={args.report}", flush=True)
    shutil.rmtree(args.temp_dir, ignore_errors=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
