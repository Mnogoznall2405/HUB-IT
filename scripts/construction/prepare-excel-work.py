"""Read-only extraction of the reviewed Svodny workbook. Never writes to a database."""
from __future__ import annotations
import argparse
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re
import sys
from uuid import NAMESPACE_URL, uuid5
from zipfile import ZipFile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'WEB-itinvent'))
from backend.models.construction_work import WorkPlan

LAYOUTS = {
    'ГПР-720-7 (101)': dict(name=3, personnel=(4, 5), plan=6, unit=7, total=8, shift=9, start=12, end=15, material=54, note=56, first=7),
    'Сооружение 200, ВНС-13': dict(name=2, personnel=(3, 4), plan=5, unit=6, total=7, shift=8, start=11, end=14, material=0, note=0, first=7),
    'ГПР-7207 (111)': dict(name=2, personnel=(0, 3), plan=4, unit=5, total=6, shift=7, start=10, end=13, material=0, note=0, first=5),
    '583-3': dict(name=3, personnel=(4, 5), plan=6, unit=7, total=8, shift=9, start=12, end=15, material=0, note=0, first=7),
}

def clean(value):
    return re.sub(r'\s+', ' ', str(value if value is not None else '')).strip()

def quantity(value, *, blank=False):
    if value is None or clean(value) in ('', '-', '—'):
        if blank: return Decimal(0)
        raise ValueError('Не указан плановый объём')
    try:
        number = Decimal(str(value).replace(' ', '').replace(',', '.'))
        if not number.is_finite() or number < 0: raise ValueError('Некорректный объём')
        return number.quantize(Decimal('.0001'))
    except InvalidOperation as exc:
        raise ValueError('Объём не является числом') from exc

def parse_date(value):
    if value is None or clean(value) in ('', '-', '—'): return None
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    for fmt in ('%d.%m.%Y', '%d/%m/%Y', '%Y-%m-%d'):
        try: return datetime.strptime(clean(value), fmt).date()
        except ValueError: pass
    raise ValueError(f'Нераспознанная дата: {clean(value)}')

def last_value_rows(path):
    # Ignore hundreds of thousands of styled empty rows, but never silently truncate data.
    result = []
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    with ZipFile(path) as archive:
        sheets = sorted((name for name in archive.namelist() if re.fullmatch(r'xl/worksheets/sheet\d+\.xml', name)), key=lambda name:int(re.search(r'(\d+)\.xml',name)[1]))
        for name in sheets:
            last = 0
            with archive.open(name) as source:
                for _, element in ET.iterparse(source, events=('end',)):
                    if element.tag == ns + 'row':
                        if any(node.text for node in element.iter() if node.tag in (ns+'v', ns+'t', ns+'f')):
                            last = int(element.attrib['r'])
                        element.clear()
            result.append(last)
    return result

def prepare(path, snapshot_date, *, unit_overrides=None):
    formulas = load_workbook(path, read_only=True, data_only=False)
    values = load_workbook(path, read_only=True, data_only=True)
    limits = last_value_rows(path)
    report = {'source': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'snapshot_date': snapshot_date.isoformat(), 'sheets': [], 'excluded_sheets': [], 'writes_performed': False}
    try:
        for sheet_index, sheet in enumerate(formulas):
            if sheet.title not in LAYOUTS:
                report['excluded_sheets'].append(sheet.title)
                continue
            layout = LAYOUTS[sheet.title]
            cached = values[sheet.title]
            output = {'sheet': sheet.title, 'last_value_row': limits[sheet_index], 'items': [], 'aggregates': [], 'blocked': [], 'warnings': [], 'auxiliary_rows': []}
            section = sheet.title
            parents = []
            building = ''
            for row, (raw, resolved) in enumerate(zip(sheet.iter_rows(max_row=limits[sheet_index], max_col=max(58,sheet.max_column)), cached.iter_rows(max_row=limits[sheet_index], max_col=max(58,sheet.max_column))), start=1):
                if row < layout['first']: continue
                ranges = {
                    'ГПР-720-7 (101)': [(7,1076),(1127,2259)],
                    'Сооружение 200, ВНС-13': [(7,118)],
                    'ГПР-7207 (111)': [(6,37)],
                    '583-3': [(7,1154)],
                }
                if not any(start <= row <= end for start,end in ranges[sheet.title]):
                    if any(c.value is not None for c in resolved): output['auxiliary_rows'].append(row)
                    continue
                cell = lambda key: resolved[layout[key]-1].value if layout[key] else None
                name = clean(cell('name'))
                if not name: continue
                number = clean(resolved[0].value)
                formula = str(raw[layout['plan']-1].value or '')
                # SUM references in planned volume identify subtotal rows in this workbook.
                aggregate = bool(re.search(r'(?i)SUM\s*\(.*[A-Z]+\d+', formula)) or (
                    formula.startswith('=') and len(re.findall(r'[A-Z]+\d+', formula)) > 1
                ) or (sheet.title == 'ГПР-7207 (111)' and row in (6,7,14,20,24,25,28,32,34))
                if aggregate:
                    label = f'{number} {name}'.strip()
                    if re.fullmatch(r'\d+(?:\.\d+)*\.?', number):
                        depth = len(number.rstrip('.').split('.'))
                        parents = parents[:depth-1] + [label]
                    elif clean(cell('unit')).startswith('усл'):
                        building = name
                        parents = []
                    else:
                        parents = parents[:1] + [label]
                    section = ' / '.join(part for part in (sheet.title, building, *parents) if part)
                    output['aggregates'].append({'row': row, 'name': name, 'number': number, 'formula': formula})
                    continue
                if not clean(cell('unit')) and cell('plan') is None and cell('total') is None:
                    section = f'{sheet.title} / {number} {name}'.strip()
                    output['aggregates'].append({'row': row, 'name': name, 'number': number, 'formula': None})
                    continue
                try:
                    notes = [clean(cell('note')), f'Источник: {path.name}; лист {sheet.title}; строка {row}.']
                    notes.append(f'Раздел исходника: {section}.')
                    display_section = section if len(section) <= 200 else section[:180] + '… [' + hashlib.sha256(section.encode()).hexdigest()[:8] + ']'
                    for label, column in zip(('ИТР в исходном срезе', 'Монтажники в исходном срезе'), layout['personnel']):
                        if column and resolved[column-1].value is not None:
                            notes.append(f'{label}: {clean(resolved[column-1].value)}.')
                    if cell('shift') is not None:
                        notes.append(f'За смену в исходном срезе: {clean(cell("shift"))}; входит в накопленный объём, повторно не прибавляется.')
                    unit = (unit_overrides or {}).get((sheet.title, row), clean(cell('unit')))
                    if not clean(cell('unit')) and unit:
                        notes.append('Единица измерения в исходнике отсутствует; требуется уточнение.')
                    plan = dict(section=display_section, name=f'{number} {name}'.strip() if number else name, unit=unit,
                                planned_quantity=quantity(cell('plan')), initial_quantity=quantity(cell('total'), blank=True),
                                initial_date=snapshot_date, weight=None, material_comment=clean(cell('material')),
                                production_comment='\n'.join(note for note in notes if note), sort_order=sheet_index*10000+row)
                    for prefix, column in (('start',layout['start']), ('end',layout['end'])):
                        for offset, mode in enumerate(('planned','revised','actual')):
                            original = resolved[column+offset-1].value
                            try:
                                parsed = parse_date(original)
                                if parsed and mode == 'actual' and parsed > date.today(): raise ValueError('Будущая фактическая дата')
                                plan[f'{mode}_{prefix}'] = parsed
                            except ValueError:
                                plan[f'{mode}_{prefix}'] = None
                                notes.append(f'Исходная дата {mode}_{prefix}: {clean(original)}; требует проверки.')
                                output['warnings'].append({'row':row,'field':f'{mode}_{prefix}','value':str(original),'reason':'Некорректная дата сохранена в примечании'})
                    for mode in ('planned','revised','actual'):
                        start,end = plan.get(f'{mode}_start'),plan.get(f'{mode}_end')
                        if start and end and start > end:
                            notes.append(f'В исходнике {mode}: начало {start}, окончание {end}. Даты противоречат друг другу, требуют проверки.')
                            output['warnings'].append({'row':row,'field':mode,'reason':'Окончание раньше начала; исходные даты сохранены в примечании'})
                            plan[f'{mode}_start'] = plan[f'{mode}_end'] = None
                    plan['production_comment'] = '\n'.join(note for note in notes if note)
                    validated = WorkPlan.model_validate(plan)
                    output['items'].append({'id': str(uuid5(NAMESPACE_URL, f'hubit-construction-svodny/{sheet.title}/{row}')), 'source_row': row, 'source_block': '583/2' if sheet.title == 'ГПР-720-7 (101)' and row >= 1127 else '583/3' if sheet.title == '583-3' else '720/7', 'plan': validated.model_dump(mode='json')})
                except (ValueError, TypeError) as exc:
                    output['blocked'].append({'row': row, 'name': name, 'reason': str(exc), 'values': [str(c.value) if c.value is not None else None for c in resolved[:17]]})
            report['sheets'].append(output)
    finally:
        formulas.close(); values.close()
    return report

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--snapshot-date', type=date.fromisoformat, default=date.today())
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    result = prepare(args.workbook, args.snapshot_date)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'sheets': [{'sheet': s['sheet'], 'last_value_row': s['last_value_row'], 'items': len(s['items']), 'aggregates': len(s['aggregates']), 'blocked': len(s['blocked'])} for s in result['sheets']], 'output': str(args.output), 'writes_performed': False}, ensure_ascii=False))
