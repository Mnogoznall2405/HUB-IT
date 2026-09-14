"""Compile the first sheet's reviewed arithmetic to live-work coefficients (read-only)."""
import argparse
import ast
from collections import defaultdict
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import re
import sys
from uuid import NAMESPACE_URL, uuid5

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'WEB-itinvent'))
from backend.services.construction_work_calculation import CalculationProfile, calculate

SHEET = 'ГПР-720-7 (101)'


def compile_profile(workbook, manifest):
    source = load_workbook(workbook, read_only=True, data_only=False)
    cached = load_workbook(workbook, read_only=True, data_only=True)
    try:
        sheet = source[SHEET]
        cells = {cell.coordinate: cell.value for row in sheet.iter_rows(max_row=1076, max_col=11) for cell in row if cell.value is not None}
        values = {cell.coordinate: cell.value for row in cached[SHEET].iter_rows(max_row=1076, max_col=11) for cell in row if cell.value is not None}
    finally:
        source.close(); cached.close()
    rows = {row['source_row']: row for sheet in manifest['sheets'] if sheet['sheet'] == SHEET for row in sheet['items'] if row['source_block'] == '720/7'}
    # The missing unit does not affect arithmetic. Activation still requires the
    # native row to exist, so it cannot silently become a frozen baseline value.
    ids = {row: data['id'] for row, data in rows.items()}
    ids[849] = str(uuid5(NAMESPACE_URL, f'hubit-construction-svodny/{SHEET}/849'))
    notes = []
    active = set()

    def combine(left, right, factor=Decimal(1)):
        terms = defaultdict(Decimal, left[1])
        for key, coefficient in right[1].items(): terms[key] += factor * coefficient
        return left[0] + factor * right[0], dict(terms)

    def scale(expression, factor):
        return expression[0] * factor, {key: coefficient * factor for key, coefficient in expression[1].items()}

    def reference(ref):
        match = re.fullmatch(r'([FH])(\d+)', ref)
        if not match: raise ValueError(f'Unreviewed reference: {ref}')
        column, row = match[1], int(match[2])
        if row in ids:
            return Decimal(0), {(ids[row], 'planned_quantity' if column == 'F' else 'total_quantity'): Decimal(1)}
        if ref in active: raise ValueError('Cyclic formula')
        value = cells.get(ref)
        if value is None: return Decimal(0), {}
        if isinstance(value, (int, float)):
            if value:
                notes.append(f'{ref}: в исходнике задан итог {value} вручную; сохранён в расчёте как постоянное значение.')
            return Decimal(str(value)), {}
        if not isinstance(value, str) or not value.startswith('='): raise ValueError(f'Non-numeric aggregate: {ref}')
        active.add(ref)
        try:
            expression = re.sub(r'([FH])(\d+):\1(\d+)', lambda m: ','.join(f'{m[1]}{n}' for n in range(int(m[2]), int(m[3]) + 1)), value[1:])
            return visit(ast.parse(expression, mode='eval').body)
        finally:
            active.remove(ref)

    def visit(node):
        if isinstance(node, ast.Name): return reference(node.id)
        if isinstance(node, ast.Constant) and type(node.value) in (int, float): return Decimal(str(node.value)), {}
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == 'SUM' and not node.keywords:
            value = Decimal(0), {}
            for argument in node.args: value = combine(value, visit(argument))
            return value
        if isinstance(node, ast.BinOp):
            left, right = visit(node.left), visit(node.right)
            if isinstance(node.op, ast.Add): return combine(left, right)
            if isinstance(node.op, ast.Sub): return combine(left, right, Decimal(-1))
            if isinstance(node.op, ast.Mult):
                if not right[1]: return scale(left, right[0])
                if not left[1]: return scale(right, left[0])
        raise ValueError(f'Unsupported arithmetic: {ast.dump(node)}')

    def expression(ref):
        constant, terms = reference(ref)
        return {'constant': str(constant), 'terms': [{'work_id': key[0], 'field': key[1], 'coefficient': str(coefficient)} for key, coefficient in sorted(terms.items()) if coefficient]}

    top = re.findall(r'K\d+', cells['K1'])
    if len(top) != 24 or cells['K1'] != '=((SUM(' + ','.join(top) + '))/2400)*100':
        raise ValueError('Top-level formula changed; review required')
    sections = []
    for ref in top:
        row = ref[1:]
        if cells[ref] != f'=H{row}/F{row}': raise ValueError('Section ratio changed')
        sections.append({'name': re.sub(r'\s+', ' ', str(cells[f'C{row}'])).strip(), 'numerator': expression(f'H{row}'), 'denominator': expression(f'F{row}')})
    profile = CalculationProfile.model_validate({
        'source_sheet': SHEET, 'source_sha256': hashlib.sha256(workbook.read_bytes()).hexdigest(),
        'group_ref': '71a8cc38-f313-11ed-bb4c-9cdc71d673ac', 'baseline_date': manifest['snapshot_date'],
        'work_ids': list(ids.values()), 'sections': sections, 'notes': sorted(set(notes)),
    })
    items = [{'id': work_id, 'group_ref': profile.group_ref, 'plan': {'planned_quantity': values.get(f'F{row}', 0), 'archived': False}, 'total_quantity': values.get(f'H{row}', 0)} for row, work_id in ids.items()]
    actual = calculate(profile, items, profile.baseline_date)
    expected = float(values['K1']) * 100
    for section, ref in zip(actual['calculation_sections'], top):
        if abs(section['percent'] - float(values[ref]) * 100) > 0.000001: raise ValueError(f'Formula and cached section disagree: {ref}')
    if abs(actual['percent'] - expected) > 0.000001: raise ValueError('Source total does not reconcile')
    return profile.model_dump(mode='json'), {'percent': actual['percent'], 'sections': len(sections), 'required_works': len(ids), 'notes': profile.notes}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('manifest', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    profile, report = compile_profile(args.workbook, json.loads(args.manifest.read_text(encoding='utf-8')))
    args.output.write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False))
