from __future__ import annotations

import argparse
import json
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import pyotp
import requests
from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


DEFAULT_API_ORIGIN = "http://127.0.0.1:8001"
DEFAULT_CHAT_API_ORIGIN = "http://127.0.0.1:8002"
TRUSTED_BROWSER_ORIGIN = "https://hubit.zsgp.ru"
ACCOUNT_PREFIX = "marketing_demo_"


class CaptureApiError(RuntimeError):
    pass


def _project_root() -> Path:
    return Path(__file__).resolve().parents[2]


def _load_json(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise SystemExit(f"Некорректный JSON: {path}")
    return payload


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _require_safe_manifest(manifest_path: Path, manifest: dict[str, Any]) -> None:
    state_root = (_project_root() / ".codex_tmp" / "about-capture").resolve()
    if state_root not in manifest_path.resolve().parents:
        raise SystemExit(f"Manifest должен находиться внутри {state_root}")
    if manifest.get("schema_version") != 1 or manifest.get("kind") != "hub-about-capture":
        raise SystemExit("Это не manifest тестового контура HUB About")
    for user in manifest.get("users") or []:
        if not str(user.get("username") or "").startswith(ACCOUNT_PREFIX):
            raise SystemExit("Manifest содержит учётку без безопасного префикса")


def _response_json(response: requests.Response, action: str) -> dict[str, Any]:
    try:
        payload = response.json() if response.content else {}
    except ValueError:
        payload = {}
    if not response.ok:
        detail = payload.get("detail") if isinstance(payload, dict) else None
        raise CaptureApiError(f"{action}: HTTP {response.status_code}: {detail or 'без описания'}")
    if not isinstance(payload, dict):
        raise CaptureApiError(f"{action}: сервер вернул неожиданный ответ")
    return payload


class HubSession:
    def __init__(
        self,
        api_origin: str,
        chat_api_origin: str,
        username: str,
        password: str,
        totp_secret: str = "",
    ):
        self.api_origin = api_origin.rstrip("/")
        self.chat_api_origin = chat_api_origin.rstrip("/")
        self.username = username
        self.password = password
        self.totp_secret = totp_secret
        self.http = requests.Session()
        self.http.headers.update(
            {
                "Accept": "application/json",
                "X-Client-Device-Id": f"hub-about-capture-{username[-20:]}",
            }
        )
        self.backup_codes: list[str] = []

    def _url(self, path: str) -> str:
        origin = self.chat_api_origin if path.startswith("/chat/") else self.api_origin
        return f"{origin}/api/v1{path}"

    def post(self, path: str, *, json_body: dict[str, Any] | None = None, **kwargs) -> dict[str, Any]:
        response = self.http.post(self._url(path), json=json_body, timeout=30, **kwargs)
        return _response_json(response, f"POST {path}")

    def get(self, path: str) -> dict[str, Any]:
        response = self.http.get(self._url(path), timeout=30)
        return _response_json(response, f"GET {path}")

    def delete(self, path: str) -> dict[str, Any]:
        response = self.http.delete(self._url(path), timeout=30)
        if response.status_code in {204, 404}:
            return {}
        return _response_json(response, f"DELETE {path}")

    def login(self) -> None:
        login = self.post(
            "/auth/login",
            json_body={"username": self.username, "password": self.password},
        )
        status = str(login.get("status") or "")
        if status == "2fa_setup_required":
            setup = self.post(
                "/auth/enable-2fa",
                json_body={"login_challenge_id": login["login_challenge_id"]},
            )
            self.totp_secret = str(setup["manual_entry_key"]).replace(" ", "")
            authenticated = self.post(
                "/auth/verify-2fa",
                json_body={
                    "login_challenge_id": login["login_challenge_id"],
                    "totp_code": pyotp.TOTP(self.totp_secret).now(),
                },
            )
            self.backup_codes = [str(item) for item in authenticated.get("backup_codes") or []]
        elif status == "2fa_required":
            if not self.totp_secret:
                raise CaptureApiError(f"Для {self.username} отсутствует TOTP secret в credentials.json")
            authenticated = self.post(
                "/auth/verify-2fa-login",
                json_body={
                    "login_challenge_id": login["login_challenge_id"],
                    "totp_code": pyotp.TOTP(self.totp_secret).now(),
                },
            )
        elif status == "authenticated":
            authenticated = login
        else:
            raise CaptureApiError(f"Неожиданный статус входа {self.username}: {status}")

        access_token = str(authenticated.get("access_token") or "")
        if not access_token:
            access_token = next(
                (
                    str(cookie.value or "")
                    for cookie in self.http.cookies
                    if cookie.name == "itinvent_access_token" and str(cookie.value or "")
                ),
                "",
            )
        if access_token:
            self.http.headers["Authorization"] = f"Bearer {access_token}"
        else:
            raise CaptureApiError(f"После входа {self.username} сервер не выдал access token")

    def playwright_cookies(self, target_host: str) -> list[dict[str, Any]]:
        cookies: list[dict[str, Any]] = []
        for cookie in self.http.cookies:
            if cookie.name not in {"itinvent_access_token", "itinvent_refresh_token"}:
                continue
            cookies.append(
                {
                    "name": cookie.name,
                    "value": str(cookie.value or ""),
                    "domain": target_host,
                    "path": "/",
                    "expires": int(cookie.expires or -1),
                    "httpOnly": True,
                    "secure": True,
                    "sameSite": "Strict",
                }
            )
        return cookies


def _make_demo_documents(target_dir: Path) -> list[tuple[Path, str]]:
    target_dir.mkdir(parents=True, exist_ok=True)

    docx_path = target_dir / "Материалы к встрече.docx"
    document = Document()
    document.add_heading("Материалы к еженедельной встрече", level=1)
    document.add_paragraph("Демонстрационный документ. Все имена и показатели вымышлены.")
    for item in ("Сверить сроки", "Подготовить краткий статус", "Собрать вопросы команды"):
        document.add_paragraph(item, style="List Bullet")
    document.save(docx_path)

    pdf_path = target_dir / "Календарь проекта.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=A4)
    pdf.setTitle("Project calendar demo")
    pdf.drawString(72, 790, "HUB-IT demo project calendar")
    pdf.drawString(72, 760, "Fictional data for presentation screenshots")
    pdf.drawString(72, 720, "Monday: planning  |  Wednesday: review  |  Friday: results")
    pdf.save()

    xlsx_path = target_dir / "Сводка по этапам.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Этапы"
    sheet.append(["Этап", "Ответственный", "Статус"])
    sheet.append(["Подготовка", "Марина Орлова", "Готово"])
    sheet.append(["Проверка", "Кирилл Волков", "В работе"])
    sheet.append(["Публикация", "Антон Лебедев", "Запланировано"])
    workbook.save(xlsx_path)

    return [
        (docx_path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        (pdf_path, "application/pdf"),
        (xlsx_path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ]


def _resource_id(payload: dict[str, Any], *keys: str) -> str:
    for key in keys:
        value = str(payload.get(key) or "").strip()
        if value:
            return value
    raise CaptureApiError(f"В ответе отсутствует ID ({', '.join(keys)})")


def _record_resource(
    manifest_path: Path,
    manifest: dict[str, Any],
    resource_key: str,
    payload: dict[str, Any],
) -> None:
    resources = manifest.setdefault("resources", {})
    resources.setdefault(resource_key, []).append(payload)
    _write_json(manifest_path, manifest)


def _save_credentials(credentials_path: Path, credentials: dict[str, Any]) -> None:
    _write_json(credentials_path, credentials)


def prepare(args: argparse.Namespace) -> int:
    if not args.confirm_current_hub:
        raise SystemExit("Повторите команду с --confirm-current-hub")
    manifest_path = Path(args.manifest).resolve()
    manifest = _load_json(manifest_path)
    _require_safe_manifest(manifest_path, manifest)
    credentials_path = manifest_path.with_name("credentials.json")
    credentials = _load_json(credentials_path)
    password = str(credentials.get("password") or "")
    if not password:
        raise SystemExit("В credentials.json отсутствует пароль")

    credential_users = {
        str(item.get("username") or ""): item
        for item in credentials.get("users") or []
    }
    sessions: dict[str, HubSession] = {}
    for user in manifest.get("users") or []:
        username = str(user["username"])
        credential = credential_users[username]
        session = HubSession(
            args.api_origin,
            args.chat_api_origin,
            username,
            password,
            str(credential.get("totp_secret") or ""),
        )
        session.login()
        credential["totp_secret"] = session.totp_secret
        if session.backup_codes:
            credential["backup_codes"] = session.backup_codes
        sessions[username] = session
        _save_credentials(credentials_path, credentials)

    users = list(manifest["users"])
    first, second, third, fourth = users
    owner = sessions[first["username"]]

    group = owner.post(
        "/chat/conversations/group",
        json_body={
            "title": "Подготовка к еженедельной встрече",
            "member_user_ids": [int(item["id"]) for item in users[1:]],
        },
    )
    group_id = _resource_id(group, "id", "conversation_id")
    _record_resource(
        manifest_path,
        manifest,
        "conversation_ids",
        {"id": group_id, "owner_username": first["username"], "kind": "group"},
    )

    direct = owner.post(
        "/chat/conversations/direct",
        json_body={"peer_user_id": int(second["id"])},
    )
    direct_id = _resource_id(direct, "id", "conversation_id")
    _record_resource(
        manifest_path,
        manifest,
        "conversation_ids",
        {"id": direct_id, "owner_username": first["username"], "kind": "direct"},
    )

    message_specs = (
        (first, "Коллеги, собрал план встречи. Проверьте свои блоки до 15:00."),
        (second, "Финансовую сводку обновила, прикреплю её к задаче."),
        (third, "Проверил доступ к материалам — всё открывается."),
        (fourth, "Добавила вопросы по поставкам. Можно обсуждать."),
    )
    message_ids: list[str] = []
    for index, (user, body) in enumerate(message_specs):
        payload: dict[str, Any] = {
            "body": body,
            "client_message_id": str(uuid.uuid4()),
        }
        if index == 2 and message_ids:
            payload["reply_to_message_id"] = message_ids[0]
        message = sessions[user["username"]].post(
            f"/chat/conversations/{group_id}/messages",
            json_body=payload,
        )
        message_ids.append(_resource_id(message, "id", "message_id"))

    for user, emoji in ((second, "👍"), (third, "✅")):
        sessions[user["username"]].post(
            f"/chat/conversations/{group_id}/messages/{message_ids[0]}/reactions",
            json_body={"emoji": emoji},
        )

    owner.post(
        f"/chat/conversations/{direct_id}/messages",
        json_body={
            "body": "Марина, посмотри, пожалуйста, итоговую задачу и сроки.",
            "client_message_id": str(uuid.uuid4()),
        },
    )
    sessions[second["username"]].post(
        f"/chat/conversations/{direct_id}/messages",
        json_body={
            "body": "Посмотрела. Всё верно, сводку уже добавила.",
            "client_message_id": str(uuid.uuid4()),
        },
    )

    projects = owner.get("/hub/task-projects").get("items") or []
    project = next((item for item in projects if item.get("is_active") is not False), None)
    if not project:
        raise CaptureApiError("В HUB нет доступного проекта задач")
    project_id = _resource_id(project, "id", "project_id")
    due_at = (datetime.now(timezone.utc) + timedelta(days=6)).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    task_response = owner.post(
        "/hub/tasks",
        json_body={
            "title": "Подготовить материалы к еженедельной встрече",
            "description": "Собрать статусы направлений, проверить вложения и согласовать повестку.",
            "assignee_user_id": int(second["id"]),
            "controller_user_id": 0,
            "observer_user_ids": [int(third["id"]), int(fourth["id"])],
            "due_at": due_at,
            "priority": "high",
            "project_id": project_id,
            "checklist_items": [
                {"text": "Собрать статусы подразделений", "done": True},
                {"text": "Проверить приложенные документы", "done": True},
                {"text": "Согласовать финальную повестку", "done": False},
            ],
        },
    )
    task_items = task_response.get("items") or []
    if not task_items:
        raise CaptureApiError("HUB не вернул созданную задачу")
    task_id = _resource_id(task_items[0], "id", "task_id")
    _record_resource(
        manifest_path,
        manifest,
        "task_ids",
        {"id": task_id, "owner_username": first["username"]},
    )

    discussion = owner.post(f"/hub/tasks/{task_id}/discussion")
    task_conversation_id = _resource_id(discussion, "conversation_id", "id")
    _record_resource(
        manifest_path,
        manifest,
        "conversation_ids",
        {"id": task_conversation_id, "owner_username": first["username"], "kind": "task"},
    )
    owner.post(
        f"/chat/conversations/{task_conversation_id}/messages",
        json_body={
            "body": "Открыл обсуждение задачи. Здесь фиксируем вопросы по материалам.",
            "client_message_id": str(uuid.uuid4()),
        },
    )
    sessions[second["username"]].post(
        f"/chat/conversations/{task_conversation_id}/messages",
        json_body={
            "body": "Два пункта уже готовы. Осталось согласовать финальную повестку.",
            "client_message_id": str(uuid.uuid4()),
        },
    )

    documents = _make_demo_documents(manifest_path.parent / "documents")
    for path, mime_type in documents:
        response = owner.http.post(
            owner._url(f"/my-files?file_name={requests.utils.quote(path.name)}&file_size={path.stat().st_size}&retention_days=7"),
            data=path.read_bytes(),
            headers={"Content-Type": mime_type, "Origin": TRUSTED_BROWSER_ORIGIN},
            timeout=60,
        )
        uploaded = _response_json(response, f"Загрузка {path.name}")
        file_id = _resource_id(uploaded, "id", "file_id")
        _record_resource(
            manifest_path,
            manifest,
            "file_ids",
            {"id": file_id, "owner_username": first["username"], "name": path.name},
        )

    attachment_path, attachment_mime = documents[0]
    response = owner.http.post(
        owner._url(f"/hub/tasks/{task_id}/attachments"),
        files={"file": (attachment_path.name, attachment_path.read_bytes(), attachment_mime)},
        timeout=60,
    )
    _response_json(response, "Загрузка вложения задачи")

    manifest["dataset_ready_at"] = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    _write_json(manifest_path, manifest)
    print(f"Тестовый сценарий готов: {manifest.get('batch')}")
    print("Создано: 2 чата, 1 обсуждение задачи, 1 задача, 3 файла.")
    return 0


def cleanup(args: argparse.Namespace) -> int:
    if not args.confirm_current_hub:
        raise SystemExit("Повторите команду с --confirm-current-hub")
    manifest_path = Path(args.manifest).resolve()
    manifest = _load_json(manifest_path)
    _require_safe_manifest(manifest_path, manifest)
    credentials = _load_json(manifest_path.with_name("credentials.json"))
    password = str(credentials.get("password") or "")
    credential_users = {
        str(item.get("username") or ""): item
        for item in credentials.get("users") or []
    }
    sessions: dict[str, HubSession] = {}

    def get_session(username: str) -> HubSession:
        if username not in sessions:
            credential = credential_users[username]
            session = HubSession(
                args.api_origin,
                args.chat_api_origin,
                username,
                password,
                str(credential.get("totp_secret") or ""),
            )
            session.login()
            sessions[username] = session
        return sessions[username]

    resources = manifest.get("resources") or {}
    for item in reversed(list(resources.get("file_ids") or [])):
        get_session(str(item["owner_username"])).delete(f"/my-files/{item['id']}")
    for item in reversed(list(resources.get("task_ids") or [])):
        get_session(str(item["owner_username"])).delete(f"/hub/tasks/{item['id']}")
    for item in reversed(list(resources.get("conversation_ids") or [])):
        if item.get("kind") == "task":
            continue
        get_session(str(item["owner_username"])).delete(f"/chat/conversations/{item['id']}")

    manifest["resources_cleanup_completed_at"] = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    _write_json(manifest_path, manifest)
    print("Ресурсы тестового сценария удалены по ID из manifest.")
    return 0


def export_browser_state(args: argparse.Namespace) -> int:
    if not args.confirm_current_hub:
        raise SystemExit("Повторите команду с --confirm-current-hub")
    manifest_path = Path(args.manifest).resolve()
    manifest = _load_json(manifest_path)
    _require_safe_manifest(manifest_path, manifest)
    credentials = _load_json(manifest_path.with_name("credentials.json"))
    credential = list(credentials.get("users") or [])[0]
    session = HubSession(
        args.api_origin,
        args.chat_api_origin,
        str(credential["username"]),
        str(credentials["password"]),
        str(credential.get("totp_secret") or ""),
    )
    session.login()
    cookies = session.playwright_cookies(args.target_host)
    if not any(item["name"] == "itinvent_access_token" for item in cookies):
        raise CaptureApiError("Не удалось получить cookie доступа для Playwright")
    state_path = manifest_path.with_name("browser-state.json")
    _write_json(state_path, {"cookies": cookies, "origins": []})
    print(f"Playwright storage state: {state_path}")
    return 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Создание обезличенных данных для скриншотов /about.")
    parser.add_argument("--manifest", required=True)
    parser.add_argument("--api-origin", default=DEFAULT_API_ORIGIN)
    parser.add_argument("--chat-api-origin", default=DEFAULT_CHAT_API_ORIGIN)
    parser.add_argument("--confirm-current-hub", action="store_true")
    subparsers = parser.add_subparsers(dest="command", required=True)
    subparsers.add_parser("prepare").set_defaults(handler=prepare)
    subparsers.add_parser("cleanup").set_defaults(handler=cleanup)
    browser_state_parser = subparsers.add_parser("export-browser-state")
    browser_state_parser.add_argument("--target-host", default="hubit.zsgp.ru")
    browser_state_parser.set_defaults(handler=export_browser_state)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    return int(args.handler(args))


if __name__ == "__main__":
    raise SystemExit(main())
