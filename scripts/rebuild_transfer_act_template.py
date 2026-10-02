"""Rebuild ``templates/docx_transfer_act.docx`` into the compact one-page layout.

Transforms the current template in place:

* the frame-based header (``w:framePr`` paragraphs) and the spacer paragraphs
  are replaced by normal flow paragraphs in the same visual order;
* the equipment table is narrowed to the text width, its header row repeats on
  every page and rows cannot split across pages;
* the two closing obligation paragraphs get ``keepNext``/``keepLines`` so the
  signature block never lands on a page alone;
* the signature grid gets a fixed-height handwriting row;
* static "1/1" / "IT Invent" frames are replaced by a real section footer
  ("HUB-IT" + ``Акт от {{DATE}}`` + ``Стр. PAGE из NUMPAGES``, gray 808080)
  and the bottom margin is raised so the footer never overlaps the text.

Idempotent: if the template has no ``w:framePr`` left, the script reports
"already rebuilt" and leaves the file untouched.

Run from the repository root::

    python scripts/rebuild_transfer_act_template.py
"""
from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Twips

PROJECT_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = PROJECT_ROOT / "templates" / "docx_transfer_act.docx"

# Text column width: page 11908 - left 720 - right 274.
TEXT_WIDTH = 10914
GRID_WIDTHS = [500, 1900, 2664, 2600, 1500, 1750]
SIG_GRID_WIDTHS = [4466, 1982, 4466]

FONT = '<w:rFonts w:ascii="Tahoma" w:hAnsi="Tahoma" w:cs="Tahoma"/>'

OBLIGATION_KEEPNEXT_PREFIXES = (
    "Материальные ценности проверены",
    "Настоящий акт составлен",
)


def _p(runs_xml: str, ppr_xml: str) -> object:
    ppr = f"<w:pPr>{ppr_xml}</w:pPr>" if ppr_xml else ""
    return parse_xml(f"<w:p {nsdecls('w')}>{ppr}{runs_xml}</w:p>")


def _text_run(text: str, *, sz: int = 19, bold: bool = False) -> str:
    b = "<w:b/><w:bCs/>" if bold else ""
    return (
        f"<w:r><w:rPr>{FONT}{b}"
        f'<w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/></w:rPr>'
        f'<w:t xml:space="preserve">{escape(text)}</w:t></w:r>'
    )


def _drawing_run(drawing) -> object:
    run = parse_xml(
        f"<w:r {nsdecls('w')}><w:rPr>{FONT}<w:noProof/>"
        f'<w:sz w:val="2"/><w:szCs w:val="2"/></w:rPr></w:r>'
    )
    run.append(drawing)
    return run


def _spacing(*, before: int = 0, after: int = 0) -> str:
    return (
        f'<w:spacing w:before="{before}" w:after="{after}" '
        f'w:line="240" w:lineRule="auto"/>'
    )


def _base_ppr(*, before: int = 0, after: int = 0) -> str:
    return (
        '<w:widowControl w:val="0"/><w:autoSpaceDE w:val="0"/>'
        '<w:autoSpaceDN w:val="0"/><w:adjustRightInd w:val="0"/>'
        + _spacing(before=before, after=after)
    )


def _text(p) -> str:
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def _set_trPr_flag(tr, flag: str) -> None:
    trPr = tr.find(qn("w:trPr"))
    if trPr is None:
        trPr = OxmlElement("w:trPr")
        tr.insert(0, trPr)
    if trPr.find(qn(f"w:{flag}")) is None:
        trPr.append(OxmlElement(f"w:{flag}"))


def _set_grid_and_widths(tbl, widths: list[int]) -> None:
    tblPr = tbl.find(qn("w:tblPr"))
    tblW = tblPr.find(qn("w:tblW"))
    tblW.set(qn("w:w"), str(sum(widths)))
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is not None:
        tblPr.remove(tblInd)
    grid_cols = tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))
    assert len(grid_cols) == len(widths)
    for grid_col, width in zip(grid_cols, widths):
        grid_col.set(qn("w:w"), str(width))
    for tr in tbl.findall(qn("w:tr")):
        cells = tr.findall(qn("w:tc"))
        assert len(cells) == len(widths)
        for tc, width in zip(cells, widths):
            tcW = tc.find(qn("w:tcPr")).find(qn("w:tcW"))
            tcW.set(qn("w:w"), str(width))


def _set_para_spacing(p, *, before=None, after=None) -> None:
    pPr = p.find(qn("w:pPr"))
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    if before is not None:
        spacing.set(qn("w:before"), str(before))
    if after is not None:
        spacing.set(qn("w:after"), str(after))


def _add_keep(p) -> None:
    pPr = p.find(qn("w:pPr"))
    if pPr.find(qn("w:keepLines")) is None:
        pPr.insert(0, OxmlElement("w:keepLines"))
    if pPr.find(qn("w:keepNext")) is None:
        pPr.insert(0, OxmlElement("w:keepNext"))


def _build_header(drawings) -> list:
    top_line, underline = drawings
    return [
        # верхняя горизонтальная линия
        _p_with_drawing(top_line, 121),
        # дата справа
        _p(_text_run("{{DATE}}"), _base_ppr() + '<w:jc w:val="right"/>'),
        # заголовок
        _p(
            _text_run(
                "АКТ ПЕРЕДАЧИ ВО ВРЕМЕННОЕ ПОЛЬЗОВАНИЕ ОБОРУДОВАНИЯ, "
                "ПРИНАДЛЕЖАЩЕГО КОМПАНИИ",
                bold=True,
            ),
            _base_ppr(before=80) + '<w:jc w:val="center"/>',
        ),
        _p(
            _text_run('ООО "Запсибгазпром-Газификация"', bold=True),
            _base_ppr() + '<w:jc w:val="center"/>',
        ),
        # «Компания» + название компании на одной строке через табуляцию
        _p(
            _text_run("Компания")
            + f'<w:r><w:rPr>{FONT}<w:sz w:val="19"/><w:szCs w:val="19"/></w:rPr>'
            '<w:tab/><w:t xml:space="preserve">ООО "Запсибгазпром-Газификация"</w:t></w:r>',
            _base_ppr(before=120)
            + '<w:tabs><w:tab w:val="left" w:pos="1801"/></w:tabs>'
            + '<w:ind w:left="800"/>',
        ),
        _p(
            _text_run("предоставляет сотруднику"),
            _base_ppr(before=60) + '<w:ind w:left="121"/>',
        ),
        _p(
            _text_run("{{TO_EMPLOYEE}}"),
            _base_ppr(before=60) + '<w:ind w:left="300"/>',
        ),
        # линия под ФИО сотрудника
        _p_with_drawing(underline, 361),
        _p(
            _text_run("во временное пользование следующее оборудование:"),
            _base_ppr(before=80) + '<w:ind w:left="121"/>',
        ),
    ]


def _p_with_drawing(drawing, ind_left: int):
    p = _p("", _base_ppr() + f'<w:ind w:left="{ind_left}"/>')
    p.append(_drawing_run(drawing))
    return p


def _footer_field_runs(instr: str, result: str) -> str:
    rpr = f"<w:rPr>{FONT}<w:color w:val=\"808080\"/><w:sz w:val=\"16\"/><w:szCs w:val=\"16\"/></w:rPr>"
    return (
        f'<w:r>{rpr}<w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r>{rpr}<w:instrText xml:space="preserve"> {instr} </w:instrText></w:r>'
        f'<w:r>{rpr}<w:fldChar w:fldCharType="separate"/></w:r>'
        f'<w:r>{rpr}<w:t>{escape(result)}</w:t></w:r>'
        f'<w:r>{rpr}<w:fldChar w:fldCharType="end"/></w:r>'
    )


def _footer_text_run(text: str) -> str:
    return (
        f'<w:r><w:rPr>{FONT}<w:color w:val="808080"/>'
        f'<w:sz w:val="16"/><w:szCs w:val="16"/></w:rPr>'
        f'<w:t xml:space="preserve">{escape(text)}</w:t></w:r>'
    )


def _build_footer(doc) -> None:
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False

    # Пустой абзац, созданный по умолчанию, — минимальной высоты над таблицей.
    empty_p = footer.paragraphs[0]._p
    parse = parse_xml(
        f"<w:pPr {nsdecls('w')}>"
        '<w:spacing w:before="0" w:after="0" w:line="20" w:lineRule="exact"/>'
        f'<w:rPr>{FONT}<w:sz w:val="2"/><w:szCs w:val="2"/></w:rPr>'
        "</w:pPr>"
    )
    empty_p.insert(0, parse)

    table = footer.add_table(1, 3, Twips(TEXT_WIDTH))
    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    tblPr.append(
        parse_xml(
            f"<w:tblBorders {nsdecls('w')}>"
            '<w:top w:val="nil"/><w:left w:val="nil"/><w:bottom w:val="nil"/>'
            '<w:right w:val="nil"/><w:insideH w:val="nil"/><w:insideV w:val="nil"/>'
            "</w:tblBorders>"
        )
    )
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    widths = [2500, 5414, 3000]
    grid_cols = tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))
    for grid_col, width in zip(grid_cols, widths):
        grid_col.set(qn("w:w"), str(width))
    row = tbl.findall(qn("w:tr"))[0]
    cells = row.findall(qn("w:tc"))
    contents = [
        (_footer_text_run("HUB-IT"), "left"),
        (_footer_text_run("Акт от {{DATE}}"), "center"),
        (
            _footer_text_run("Стр. ")
            + _footer_field_runs("PAGE", "1")
            + _footer_text_run(" из ")
            + _footer_field_runs("NUMPAGES", "1"),
            "right",
        ),
    ]
    for tc, width, (runs_xml, jc) in zip(cells, widths, contents):
        tc.find(qn("w:tcPr")).find(qn("w:tcW")).set(qn("w:w"), str(width))
        old_p = tc.find(qn("w:p"))
        tc.remove(old_p)
        tc.append(
            _p(
                runs_xml,
                '<w:spacing w:before="0" w:after="0" w:line="240" '
                f'w:lineRule="auto"/><w:jc w:val="{jc}"/>',
            )
        )


def rebuild(template_path: Path = TEMPLATE_PATH) -> None:
    doc = Document(str(template_path))
    body = doc.element.body

    if body.find(f".//{qn('w:framePr')}") is None:
        print("already rebuilt: no w:framePr found, template left untouched")
        return

    children = list(body)
    sectPr = children[-1]
    assert sectPr.tag == qn("w:sectPr")
    top_tables = [el for el in children if el.tag == qn("w:tbl")]
    tbl_equipment, tbl_signature = top_tables

    # --- сохранить картинки-линии до удаления рамок ---
    drawings = []
    for p in children:
        if p.tag != qn("w:p"):
            continue
        for drawing in p.iter(qn("w:drawing")):
            drawings.append(drawing)
    assert len(drawings) == 2, f"expected 2 line images, got {len(drawings)}"

    # --- шапка: удалить все абзацы до первой таблицы (рамки + распорки) ---
    for el in children[: children.index(tbl_equipment)]:
        body.remove(el)
    for header_p in _build_header(drawings):
        tbl_equipment.addprevious(header_p)

    # --- таблица техники: ширины, повтор шапки, cantSplit ---
    _set_grid_and_widths(tbl_equipment, GRID_WIDTHS)
    rows = tbl_equipment.findall(qn("w:tr"))
    _set_trPr_flag(rows[0], "cantSplit")
    _set_trPr_flag(rows[0], "tblHeader")
    for tr in rows[1:]:
        _set_trPr_flag(tr, "cantSplit")

    # --- абзацы между таблицей техники и блоком подписей ---
    closing_paras = []
    for p in body.findall(qn("w:p")):
        text = _text(p)
        if any(text.startswith(prefix) for prefix in OBLIGATION_KEEPNEXT_PREFIXES):
            closing_paras.append(p)
    assert len(closing_paras) == 2
    for p in closing_paras:
        _add_keep(p)
    _set_para_spacing(closing_paras[-1], after=300)

    # пустые абзацы между таблицей техники и блоком подписей: распорка после
    # таблицы, пустой абзац после «Сотрудник принимает…», старая пустая рамка
    # и две распорки перед подписями
    removable = []
    seen_equipment = False
    for el in list(body):
        if el is tbl_equipment:
            seen_equipment = True
            continue
        if not seen_equipment:
            continue
        if el is tbl_signature:
            break
        if el.tag == qn("w:p") and _text(el) == "":
            removable.append(el)
    assert len(removable) == 5, f"expected 5 spacer paragraphs, got {len(removable)}"
    for el in removable:
        body.remove(el)
    # отступы вместо распорок
    for p in body.findall(qn("w:p")):
        if _text(p).startswith("Сотрудник принимает на себя следующие обязательства"):
            _set_para_spacing(p, before=120, after=60)
            break

    # --- блок подписей: ширины и компактная строка под подпись ---
    _set_grid_and_widths(tbl_signature, [TEXT_WIDTH])
    inner_tbl = tbl_signature.find(f".//{qn('w:tbl')}")
    _set_grid_and_widths(inner_tbl, SIG_GRID_WIDTHS)
    inner_rows = inner_tbl.findall(qn("w:tr"))
    name_row_height = inner_rows[2].find(qn("w:trPr")).find(qn("w:trHeight"))
    name_row_height.set(qn("w:val"), "440")
    for p in inner_rows[1].iter(qn("w:p")):
        _set_para_spacing(p, before=80, after=40)

    # --- обязательный финальный абзац: минимальная высота ---
    final_p = tbl_signature.getnext()
    assert final_p.tag == qn("w:p")
    _set_para_spacing(final_p, before=0, after=0)
    spacing = final_p.find(qn("w:pPr")).find(qn("w:spacing"))
    spacing.set(qn("w:line"), "20")
    spacing.set(qn("w:lineRule"), "exact")
    rPr = final_p.find(qn("w:pPr")).find(qn("w:rPr"))
    for tag in ("w:sz", "w:szCs"):
        el = rPr.find(qn(tag))
        el.set(qn("w:val"), "2")

    # --- поля страницы: место под колонтитул ---
    pgMar = sectPr.find(qn("w:pgMar"))
    pgMar.set(qn("w:bottom"), "640")
    pgMar.set(qn("w:footer"), "288")

    _build_footer(doc)

    # --- стиль таблицы: новые строки данных в Tahoma 10 pt, как шапка ---
    styles_el = doc.styles.element
    for style in styles_el.findall(qn("w:style")):
        if style.get(qn("w:styleId")) == "a3":
            tblPr_el = style.find(qn("w:tblPr"))
            tblPr_el.addprevious(
                parse_xml(
                    f"<w:rPr {nsdecls('w')}>{FONT}"
                    '<w:sz w:val="20"/><w:szCs w:val="20"/></w:rPr>'
                )
            )
            break

    doc.save(str(template_path))
    print(f"rebuilt: {template_path}")


if __name__ == "__main__":
    rebuild()
