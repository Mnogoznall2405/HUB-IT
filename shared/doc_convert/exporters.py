from __future__ import annotations

import io
import re
from typing import Any

from shared.doc_convert.models import ExportFormat, SignatureCrop

DOC_CONVERT_FORMATS: tuple[ExportFormat, ...] = ("docx", "txt", "md", "pdf", "xlsx")

_FORMAT_ALIASES = {
    "word": "docx",
    "document": "docx",
    "doc": "docx",
    "text": "txt",
    "plaintext": "txt",
    "markdown": "md",
    "excel": "xlsx",
    "xls": "xlsx",
    "spreadsheet": "xlsx",
}


def normalize_export_format(value: object) -> ExportFormat:
    text = str(value or "").strip().lower().lstrip(".")
    text = _FORMAT_ALIASES.get(text, text)
    if text not in DOC_CONVERT_FORMATS:
        raise ValueError(f"Unsupported export format: {value}")
    return text  # type: ignore[return-value]


def _safe_stem(file_name: str) -> str:
    stem = re.sub(r"[^\w\-а-яА-ЯёЁ]+", "_", str(file_name or "document").strip(), flags=re.UNICODE)
    stem = stem.strip("._") or "document"
    return stem[:120]


def markdown_to_plain_text(markdown: str) -> str:
    lines: list[str] = []
    for raw in str(markdown or "").replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw.rstrip()
        if line.startswith("#"):
            line = line.lstrip("#").strip()
        line = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        line = re.sub(r"`([^`]+)`", r"\1", line)
        lines.append(line)
    return "\n".join(lines).strip() + ("\n" if markdown else "")


def _parse_markdown_tables(markdown: str) -> list[list[list[str]]]:
    tables: list[list[list[str]]] = []
    lines = str(markdown or "").splitlines()
    index = 0

    def split_row(line: str) -> list[str]:
        text = line.strip().strip("|")
        return [cell.strip() for cell in text.split("|")]

    while index < len(lines):
        line = lines[index].strip()
        if "|" not in line:
            index += 1
            continue
        if index + 1 >= len(lines):
            break
        sep = lines[index + 1].strip()
        if not re.match(r"^\|?\s*:?-{3,}", sep):
            index += 1
            continue
        rows = [split_row(line)]
        index += 2
        while index < len(lines) and "|" in lines[index]:
            rows.append(split_row(lines[index]))
            index += 1
        if rows:
            tables.append(rows)
    return tables


def _markdown_blocks(markdown: str) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    lines = str(markdown or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    index = 0
    paragraph: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph
        text = "\n".join(paragraph).strip()
        paragraph = []
        if text:
            blocks.append({"type": "paragraph", "text": text})

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("#"):
            flush_paragraph()
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            blocks.append({"type": "heading", "level": max(1, min(level, 4)), "text": text})
            index += 1
            continue
        if "|" in stripped and index + 1 < len(lines) and re.match(r"^\|?\s*:?-{3,}", lines[index + 1].strip()):
            flush_paragraph()
            table_lines = [stripped]
            index += 1
            table_lines.append(lines[index].strip())
            index += 1
            while index < len(lines) and "|" in lines[index]:
                table_lines.append(lines[index].strip())
                index += 1
            table_md = "\n".join(table_lines)
            tables = _parse_markdown_tables(table_md)
            if tables:
                blocks.append({"type": "table", "rows": tables[0]})
            continue
        if re.match(r"^[-*+]\s+", stripped) or re.match(r"^\d+\.\s+", stripped):
            flush_paragraph()
            items: list[str] = []
            while index < len(lines):
                current = lines[index].strip()
                bullet = re.match(r"^[-*+]\s+(.*)$", current)
                numbered = re.match(r"^\d+\.\s+(.*)$", current)
                if not bullet and not numbered:
                    break
                items.append((bullet or numbered).group(1).strip())
                index += 1
            blocks.append({"type": "list", "items": items})
            continue
        paragraph.append(line)
        index += 1
    flush_paragraph()
    return blocks


def _apply_inline_markdown_runs(paragraph: Any, text: str, *, force_bold: bool = False) -> None:
    """Render **bold** and __underline__ markers into python-docx runs."""
    raw = str(text or "")
    if not raw:
        paragraph.add_run("")
        return
    # Split on bold and underline tokens; process nested simply left-to-right.
    parts = re.split(r"(\*\*[^*]+\*\*|__[^_]+__)", raw)
    for part in parts:
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("__") and part.endswith("__") and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.underline = True
            if force_bold:
                run.bold = True
        elif part:
            run = paragraph.add_run(part)
            if force_bold:
                run.bold = True


def _docx_helpers():
    from docx import Document  # type: ignore
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    from docx.shared import Inches, Pt  # type: ignore

    def set_table_full_width(table: Any) -> None:
        try:
            table.autofit = True
            tbl = table._tbl
            tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
            if tbl.tblPr is None:
                tbl.insert(0, tbl_pr)
            width = OxmlElement("w:tblW")
            width.set(qn("w:w"), "5000")
            width.set(qn("w:type"), "pct")
            for child in list(tbl_pr):
                if child.tag == qn("w:tblW"):
                    tbl_pr.remove(child)
            tbl_pr.append(width)
        except Exception:
            pass

    def set_col_widths(table: Any, ratios: list[float]) -> None:
        try:
            from docx.shared import Twips  # type: ignore

            total = 9000
            widths = [Twips(int(total * ratio)) for ratio in ratios]
            for row in table.rows:
                for index, cell in enumerate(row.cells):
                    if index < len(widths):
                        cell.width = widths[index]
        except Exception:
            pass

    def reset_paragraph(paragraph: Any) -> None:
        for run in list(paragraph.runs):
            try:
                run._element.getparent().remove(run._element)
            except Exception:
                run.text = ""

    def write_cell(cell: Any, text: str, *, bold: bool = False) -> None:
        paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        reset_paragraph(paragraph)
        _apply_inline_markdown_runs(paragraph, str(text or ""))
        if bold:
            for run in paragraph.runs:
                run.bold = True
        for run in paragraph.runs:
            run.font.size = Pt(11)

    return Document, WD_ALIGN_PARAGRAPH, Inches, Pt, set_table_full_width, set_col_widths, reset_paragraph, write_cell


def _picture_size_inches(png_bytes: bytes, *, max_width: float = 1.9, max_height: float = 0.75) -> tuple[float, float]:
    try:
        from PIL import Image  # type: ignore

        with Image.open(io.BytesIO(png_bytes)) as image:
            width_px, height_px = image.size
    except Exception:
        return max_width, max_height * 0.6
    if width_px < 1 or height_px < 1:
        return max_width, max_height * 0.6
    # Assume ~150 dpi crop from zoomed page render.
    width_in = width_px / 150.0
    height_in = height_px / 150.0
    scale = min(max_width / width_in, max_height / height_in, 1.0)
    return max(0.4, width_in * scale), max(0.2, height_in * scale)


def _float_inline_picture_over_line(run: Any, *, lift_inches: float = 0.28) -> None:
    """Convert the run's inline picture into a floating overlay above the text line."""
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    from docx.shared import Inches  # type: ignore

    drawing = None
    for child in run._element:
        if child.tag == qn("w:drawing"):
            drawing = child
    if drawing is None:
        return
    inline = None
    for child in drawing:
        if child.tag == qn("wp:inline"):
            inline = child
            break
    if inline is None:
        return

    extent = inline.find(qn("wp:extent"))
    doc_pr = inline.find(qn("wp:docPr"))
    cnv = inline.find(qn("wp:cNvGraphicFramePr"))
    graphic = inline.find(qn("a:graphic"))
    if extent is None or doc_pr is None or graphic is None:
        return

    cx = extent.get("cx") or "0"
    cy = extent.get("cy") or "0"
    lift_emu = str(int(Inches(lift_inches)))

    anchor = OxmlElement("wp:anchor")
    anchor.set("distT", "0")
    anchor.set("distB", "0")
    anchor.set("distL", "0")
    anchor.set("distR", "0")
    anchor.set("simplePos", "0")
    anchor.set("relativeHeight", "251658240")
    anchor.set("behindDoc", "0")  # in front of text / line
    anchor.set("locked", "0")
    anchor.set("layoutInCell", "1")
    anchor.set("allowOverlap", "1")

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)

    pos_h = OxmlElement("wp:positionH")
    pos_h.set("relativeFrom", "column")
    align_h = OxmlElement("wp:align")
    align_h.text = "center"
    pos_h.append(align_h)
    anchor.append(pos_h)

    pos_v = OxmlElement("wp:positionV")
    pos_v.set("relativeFrom", "paragraph")
    # Negative offset lifts the ink onto/over the underscore.
    pos_offset = OxmlElement("wp:posOffset")
    pos_offset.text = f"-{lift_emu}"
    pos_v.append(pos_offset)
    anchor.append(pos_v)

    extent_copy = OxmlElement("wp:extent")
    extent_copy.set("cx", cx)
    extent_copy.set("cy", cy)
    anchor.append(extent_copy)

    effect = OxmlElement("wp:effectExtent")
    effect.set("l", "0")
    effect.set("t", "0")
    effect.set("r", "0")
    effect.set("b", "0")
    anchor.append(effect)

    wrap = OxmlElement("wp:wrapNone")
    anchor.append(wrap)

    anchor.append(doc_pr)
    if cnv is not None:
        anchor.append(cnv)
    anchor.append(graphic)

    drawing.remove(inline)
    drawing.append(anchor)


def _add_signature_overlay_picture(paragraph: Any, png_bytes: bytes) -> bool:
    from docx.shared import Inches  # type: ignore

    width_in, height_in = _picture_size_inches(png_bytes)
    try:
        run = paragraph.add_run()
        run.add_picture(io.BytesIO(png_bytes), width=Inches(width_in), height=Inches(height_in))
        _float_inline_picture_over_line(run, lift_inches=max(0.18, height_in * 0.55))
        return True
    except Exception:
        return False


def _export_docx_from_structure(
    structure_pages: list[dict[str, Any]],
    *,
    signatures: list[SignatureCrop] | None = None,
) -> bytes:
    from shared.doc_convert.structure import normalize_structure_payload

    try:
        (
            Document,
            WD_ALIGN_PARAGRAPH,
            Inches,
            Pt,
            set_table_full_width,
            set_col_widths,
            reset_paragraph,
            write_cell,
        ) = _docx_helpers()
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("python-docx is not installed") from exc

    align_map = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }

    def add_signature_block(left: str, right: str, crop: SignatureCrop | None) -> None:
        if not (left or right or (crop is not None and crop.png_bytes)):
            return
        document.add_paragraph("")
        sig_table = document.add_table(rows=1, cols=3)
        set_table_full_width(sig_table)
        set_col_widths(sig_table, [0.38, 0.28, 0.34])
        try:
            sig_table.style = None
        except Exception:
            pass
        write_cell(sig_table.rows[0].cells[0], left or (crop.left_text if crop else ""))
        middle = sig_table.rows[0].cells[1]
        reset_paragraph(middle.paragraphs[0])
        line_paragraph = middle.paragraphs[0]
        line_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        line_paragraph.add_run("_______________")
        if crop is not None and crop.png_bytes:
            if not _add_signature_overlay_picture(line_paragraph, crop.png_bytes):
                try:
                    pic_p = middle.add_paragraph()
                    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    width_in, height_in = _picture_size_inches(crop.png_bytes)
                    pic_p.add_run().add_picture(
                        io.BytesIO(crop.png_bytes),
                        width=Inches(width_in),
                        height=Inches(height_in),
                    )
                except Exception:
                    pass
        write_cell(sig_table.rows[0].cells[2], right or (crop.right_text if crop else ""))

    signatures = list(signatures or [])
    signature_by_page = {int(item.page_index): item for item in signatures}
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    for enum_index, raw_payload in enumerate(structure_pages):
        payload = normalize_structure_payload(raw_payload)
        page_index = int(payload.get("page_index") or (enum_index + 1))
        if enum_index > 0:
            document.add_page_break()

        title = str(payload.get("title") or "").strip()
        if title:
            heading = document.add_heading("", level=1)
            reset_paragraph(heading)
            run = heading.add_run(title)
            run.bold = True
            run.font.size = Pt(14)
            if bool(payload.get("title_centered", True)):
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        page_crop = signature_by_page.get(page_index)
        if page_crop is None and signatures and len(structure_pages) == 1:
            page_crop = signatures[0]

        prev_y1 = 0.0
        for block in list(payload.get("blocks") or []):
            bbox = list(block.get("bbox_norm") or [])
            if len(bbox) == 4:
                try:
                    y0, y1 = float(bbox[1]), float(bbox[3])
                    if y1 - y0 >= 0.01:
                        if prev_y1 > 0 and (y0 - prev_y1) > 0.035:
                            document.add_paragraph("")
                        prev_y1 = y1
                except Exception:
                    pass
            kind = str(block.get("type") or "").strip().lower()
            if kind == "paragraph":
                text = str(block.get("text") or "").strip()
                if not text:
                    continue
                paragraph = document.add_paragraph()
                paragraph.alignment = align_map.get(str(block.get("align") or "left"), WD_ALIGN_PARAGRAPH.LEFT)
                if bool(block.get("indent_first")) and str(block.get("align") or "left") == "left":
                    paragraph.paragraph_format.first_line_indent = Inches(0.5)
                _apply_inline_markdown_runs(paragraph, text, force_bold=bool(block.get("bold")))
                for run in paragraph.runs:
                    run.font.size = Pt(12)
                    run.font.name = "Times New Roman"
            elif kind == "list":
                marker = "—" if str(block.get("list_marker") or "dash") == "dash" else "•"
                for item in list(block.get("items") or []):
                    text = str(item or "").strip()
                    if not text:
                        continue
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.left_indent = Inches(0.25)
                    _apply_inline_markdown_runs(paragraph, f"{marker} {text}")
                    for run in paragraph.runs:
                        run.font.size = Pt(12)
                        run.font.name = "Times New Roman"
            elif kind == "table":
                rows = list(block.get("rows") or [])
                if not rows:
                    continue
                cols = max((len(row) for row in rows), default=0) or 1
                table = document.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                set_table_full_width(table)
                if cols == 2:
                    set_col_widths(table, [0.42, 0.58])
                for r_index, row in enumerate(rows):
                    for c_index in range(cols):
                        value = row[c_index] if c_index < len(row) else ""
                        write_cell(
                            table.rows[r_index].cells[c_index],
                            str(value),
                            bold=(cols == 2 and c_index == 0),
                        )
            elif kind == "signature":
                add_signature_block(
                    str(block.get("left_text") or "").strip(),
                    str(block.get("right_text") or "").strip(),
                    page_crop,
                )

    if not document.paragraphs and not document.tables:
        document.add_paragraph("")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _export_docx(markdown: str) -> bytes:
    try:
        (
            Document,
            WD_ALIGN_PARAGRAPH,
            _Inches,
            _Pt,
            set_table_full_width,
            _set_col_widths,
            reset_paragraph,
            write_cell,
        ) = _docx_helpers()
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("python-docx is not installed") from exc

    document = Document()
    for block in _markdown_blocks(markdown):
        kind = block.get("type")
        if kind == "heading":
            heading = document.add_heading("", level=int(block.get("level") or 1))
            reset_paragraph(heading)
            _apply_inline_markdown_runs(heading, str(block.get("text") or ""))
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "paragraph":
            paragraph = document.add_paragraph()
            _apply_inline_markdown_runs(paragraph, str(block.get("text") or ""))
        elif kind == "list":
            for item in list(block.get("items") or []):
                paragraph = document.add_paragraph(style="List Bullet")
                reset_paragraph(paragraph)
                _apply_inline_markdown_runs(paragraph, str(item))
        elif kind == "table":
            rows = list(block.get("rows") or [])
            if not rows:
                continue
            cols = max(len(row) for row in rows) or 1
            table = document.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            set_table_full_width(table)
            for r_index, row in enumerate(rows):
                for c_index in range(cols):
                    value = row[c_index] if c_index < len(row) else ""
                    write_cell(
                        table.rows[r_index].cells[c_index],
                        str(value),
                        bold=(cols == 2 and c_index == 0),
                    )
    if not document.paragraphs and not document.tables:
        document.add_paragraph(markdown_to_plain_text(markdown).strip() or "")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _strip_inline_md(text: object) -> str:
    value = str(text or "")
    value = re.sub(r"\*\*(.+?)\*\*", r"\1", value)
    value = re.sub(r"__(.+?)__", r"\1", value)
    value = re.sub(r"`([^`]+)`", r"\1", value)
    return value.strip()


def _sanitize_sheet_title(raw_title: object, *, index: int, used: set[str]) -> str:
    fallback = f"Лист{index}"
    title = re.sub(r'[\\/*?:\[\]]+', "_", str(raw_title or "").strip())
    title = title.strip(" .'\"") or fallback
    base = title[:31].strip(" .") or fallback[:31]
    candidate = base
    suffix = 2
    while candidate.casefold() in used:
        ending = f"_{suffix}"
        candidate = f"{base[: max(1, 31 - len(ending))].strip(' .')}{ending}"
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _xlsx_style_sheet(sheet, *, header_row: int | None = 1, bold_first_col: bool = False) -> None:
    try:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # type: ignore
        from openpyxl.utils import get_column_letter  # type: ignore
    except Exception:
        return

    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(vertical="top", wrap_text=True)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    label_font = Font(bold=True)

    max_row = int(sheet.max_row or 0)
    max_col = int(sheet.max_column or 0)
    if max_row < 1 or max_col < 1:
        return

    for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            cell.alignment = wrap
            cell.border = border
            if bold_first_col and cell.column == 1 and (header_row is None or cell.row != header_row):
                cell.font = label_font

    if header_row is not None and 1 <= header_row <= max_row:
        for cell in sheet[header_row]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.freeze_panes = f"A{header_row + 1}"
        try:
            sheet.auto_filter.ref = sheet.dimensions
        except Exception:
            pass

    for col_index in range(1, max_col + 1):
        width = 12
        for row_index in range(1, max_row + 1):
            value = sheet.cell(row=row_index, column=col_index).value
            width = max(width, min(48, len(str(value or "")) + 2))
        sheet.column_dimensions[get_column_letter(col_index)].width = width


def _normalize_table_rows(rows: list[Any]) -> list[list[str]]:
    normalized: list[list[str]] = []
    width = 0
    for row in rows:
        cells = [_strip_inline_md(cell) for cell in list(row or [])]
        if not any(cells):
            continue
        width = max(width, len(cells))
        normalized.append(cells)
    if width < 1:
        return []
    return [row + [""] * (width - len(row)) for row in normalized]


def _write_table_sheet(workbook, *, title: str, rows: list[list[str]], used_titles: set[str], index: int):
    sheet_title = _sanitize_sheet_title(title, index=index, used=used_titles)
    sheet = workbook.create_sheet(sheet_title)

    # 2-column forms → label|value with explicit header when first row is data.
    if rows and len(rows[0]) == 2:
        left, right = rows[0][0].casefold(), rows[0][1].casefold()
        looks_header = left in {"поле", "показатель", "параметр", "name", "field"} or right in {
            "значение",
            "value",
            "данные",
        }
        if not looks_header:
            sheet.cell(row=1, column=1, value="Поле")
            sheet.cell(row=1, column=2, value="Значение")
            for r_index, row in enumerate(rows, start=2):
                sheet.cell(row=r_index, column=1, value=row[0] if len(row) > 0 else "")
                sheet.cell(row=r_index, column=2, value=row[1] if len(row) > 1 else "")
            _xlsx_style_sheet(sheet, header_row=1, bold_first_col=True)
            return sheet

    for r_index, row in enumerate(rows, start=1):
        for c_index, value in enumerate(row, start=1):
            sheet.cell(row=r_index, column=c_index, value=value)
    _xlsx_style_sheet(sheet, header_row=1, bold_first_col=False)
    return sheet


def _export_xlsx_from_structure(structure_pages: list[dict[str, Any]]) -> tuple[bytes, list[str]]:
    try:
        from openpyxl import Workbook  # type: ignore
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is not installed") from exc

    from shared.doc_convert.structure import normalize_structure_payload

    warnings: list[str] = []
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles: set[str] = set()
    text_rows: list[list[str]] = [["Страница", "Тип", "Текст"]]
    table_count = 0
    multi_page = len(structure_pages) > 1

    for enum_index, raw_payload in enumerate(structure_pages):
        payload = normalize_structure_payload(raw_payload)
        page_index = int(payload.get("page_index") or (enum_index + 1))
        title = _strip_inline_md(payload.get("title"))
        if title:
            text_rows.append([str(page_index), "Заголовок", title])

        for block in list(payload.get("blocks") or []):
            kind = str(block.get("type") or "").strip().lower()
            if kind == "paragraph":
                text = _strip_inline_md(block.get("text"))
                if text:
                    text_rows.append([str(page_index), "Абзац", text])
            elif kind == "list":
                marker = "—" if str(block.get("list_marker") or "dash") == "dash" else "•"
                for item in list(block.get("items") or []):
                    text = _strip_inline_md(item)
                    if text:
                        text_rows.append([str(page_index), "Список", f"{marker} {text}"])
            elif kind == "signature":
                left = _strip_inline_md(block.get("left_text"))
                right = _strip_inline_md(block.get("right_text"))
                if left or right:
                    text_rows.append([str(page_index), "Подпись", f"{left}  ________  {right}".strip()])
            elif kind == "table":
                rows = _normalize_table_rows(list(block.get("rows") or []))
                if not rows:
                    continue
                table_count += 1
                if multi_page:
                    sheet_name = f"Стр{page_index}_Т{table_count}"
                else:
                    sheet_name = f"Таблица{table_count}"
                if title and table_count == 1 and not multi_page:
                    # Prefer human title for the first/only form table.
                    sheet_name = title[:28] or sheet_name
                _write_table_sheet(
                    workbook,
                    title=sheet_name,
                    rows=rows,
                    used_titles=used_titles,
                    index=table_count,
                )

    if len(text_rows) > 1:
        doc_sheet = workbook.create_sheet(
            _sanitize_sheet_title("Документ", index=1, used=used_titles),
            0,
        )
        for r_index, row in enumerate(text_rows, start=1):
            for c_index, value in enumerate(row, start=1):
                doc_sheet.cell(row=r_index, column=c_index, value=value)
        _xlsx_style_sheet(doc_sheet, header_row=1, bold_first_col=False)
    elif table_count == 0:
        warnings.append("В структуре нет таблиц и текста — создан пустой лист.")
        empty = workbook.create_sheet(_sanitize_sheet_title("Документ", index=1, used=used_titles))
        empty.cell(row=1, column=1, value="")
    elif table_count > 0:
        # Tables-only document is fine.
        pass

    if table_count == 0 and len(text_rows) > 1:
        warnings.append("Таблиц в документе нет — текст вынесен на лист «Документ».")

    if not workbook.worksheets:
        workbook.create_sheet("Документ")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), warnings


def _export_xlsx(markdown: str) -> tuple[bytes, list[str]]:
    try:
        from openpyxl import Workbook  # type: ignore
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is not installed") from exc

    warnings: list[str] = []
    tables = _parse_markdown_tables(markdown)
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles: set[str] = set()

    # Collect non-table text for a companion sheet.
    text_lines: list[str] = []
    for block in _markdown_blocks(markdown):
        kind = block.get("type")
        if kind == "heading":
            text = _strip_inline_md(block.get("text"))
            if text:
                text_lines.append(text)
        elif kind == "paragraph":
            text = _strip_inline_md(block.get("text"))
            if text:
                text_lines.append(text)
        elif kind == "list":
            for item in list(block.get("items") or []):
                text = _strip_inline_md(item)
                if text:
                    text_lines.append(f"• {text}")

    if tables:
        for table_index, rows in enumerate(tables, start=1):
            normalized = _normalize_table_rows(rows)
            if not normalized:
                continue
            _write_table_sheet(
                workbook,
                title=f"Таблица{table_index}",
                rows=normalized,
                used_titles=used_titles,
                index=table_index,
            )
        if text_lines:
            sheet = workbook.create_sheet(
                _sanitize_sheet_title("Документ", index=1, used=used_titles),
                0,
            )
            sheet.cell(row=1, column=1, value="Текст")
            for index, line in enumerate(text_lines, start=2):
                sheet.cell(row=index, column=1, value=line)
            _xlsx_style_sheet(sheet, header_row=1)
    else:
        warnings.append("В документе не найдено таблиц — Excel содержит текст построчно.")
        sheet = workbook.create_sheet(_sanitize_sheet_title("Документ", index=1, used=used_titles))
        sheet.cell(row=1, column=1, value="Текст")
        lines = text_lines or markdown_to_plain_text(markdown).splitlines() or [""]
        for index, line in enumerate(lines, start=2):
            sheet.cell(row=index, column=1, value=_strip_inline_md(line))
        _xlsx_style_sheet(sheet, header_row=1)

    if not workbook.worksheets:
        workbook.create_sheet("Документ")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), warnings


def _pdf_font_path() -> str | None:
    from pathlib import Path

    candidates = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/ARIAL.TTF"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return str(path)
    return None


def _export_pdf(markdown: str, *, title: str = "") -> bytes:
    try:
        from reportlab.lib.pagesizes import A4  # type: ignore
        from reportlab.pdfbase import pdfmetrics  # type: ignore
        from reportlab.pdfbase.ttfonts import TTFont  # type: ignore
        from reportlab.pdfgen import canvas  # type: ignore
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("reportlab is not installed") from exc

    font_path = _pdf_font_path()
    if not font_path:
        raise RuntimeError("PDF font unavailable")
    font_name = "DocConvertSans"
    try:
        pdfmetrics.getFont(font_name)
    except Exception:
        pdfmetrics.registerFont(TTFont(font_name, font_path))

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    y = height - 48
    text = markdown_to_plain_text(markdown)
    lines = text.splitlines() or [""]
    if title:
        lines = [str(title), ""] + lines

    def new_page() -> None:
        nonlocal y
        pdf.showPage()
        pdf.setFont(font_name, 10)
        y = height - 48

    pdf.setFont(font_name, 10)
    for raw_line in lines:
        chunks = [raw_line[i : i + 100] for i in range(0, max(1, len(raw_line)), 100)] or [""]
        for chunk in chunks:
            if y < 48:
                new_page()
            pdf.drawString(40, y, chunk)
            y -= 14
    pdf.save()
    return buffer.getvalue()


def export_markdown(
    markdown: str,
    *,
    export_format: object,
    source_name: str = "document",
    title: str = "",
    structure_pages: list[dict[str, Any]] | None = None,
    signatures: list[SignatureCrop] | None = None,
) -> dict[str, Any]:
    fmt = normalize_export_format(export_format)
    stem = _safe_stem(source_name.rsplit(".", 1)[0] if "." in source_name else source_name)
    warnings: list[str] = []
    structure_pages = [item for item in list(structure_pages or []) if isinstance(item, dict)]
    signatures = list(signatures or [])
    if fmt == "md":
        payload = (str(markdown or "").strip() + "\n").encode("utf-8")
        mime = "text/markdown"
    elif fmt == "txt":
        payload = markdown_to_plain_text(markdown).encode("utf-8")
        mime = "text/plain"
    elif fmt == "docx":
        if structure_pages:
            payload = _export_docx_from_structure(structure_pages, signatures=signatures)
        else:
            payload = _export_docx(markdown)
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif fmt == "xlsx":
        if structure_pages:
            payload, warnings = _export_xlsx_from_structure(structure_pages)
        else:
            payload, warnings = _export_xlsx(markdown)
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif fmt == "pdf":
        payload = _export_pdf(markdown, title=title or stem)
        mime = "application/pdf"
    else:  # pragma: no cover
        raise ValueError(f"Unsupported export format: {fmt}")
    return {
        "format": fmt,
        "file_name": f"{stem}.{fmt}",
        "mime_type": mime,
        "content": payload,
        "warnings": warnings,
    }
