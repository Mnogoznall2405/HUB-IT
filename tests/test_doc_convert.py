from __future__ import annotations

import io
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
WEB_ROOT = PROJECT_ROOT / "WEB-itinvent"

if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
if str(WEB_ROOT) not in sys.path:
    sys.path.insert(0, str(WEB_ROOT))


def test_merge_and_table_count():
    from shared.doc_convert.merge import count_markdown_tables, merge_page_markdown, strip_markdown_fences
    from shared.doc_convert.models import PageMarkdown

    assert strip_markdown_fences("```markdown\n# Title\n```") == "# Title"
    pages = [
        PageMarkdown(page_index=1, markdown="# Doc\n\nHeader Co\n\nHello", source="text", char_count=10),
        PageMarkdown(page_index=2, markdown="Header Co\n\n| A | B |\n| --- | --- |\n| 1 | 2 |", source="vision", char_count=20),
    ]
    merged = merge_page_markdown(pages)
    assert "Hello" in merged
    assert count_markdown_tables(merged) == 1


def test_export_markdown_formats_roundtrip():
    from shared.doc_convert.exporters import export_markdown

    markdown = "# Title\n\nParagraph\n\n- one\n- two\n\n| A | B |\n| --- | --- |\n| 1 | 2 |\n"
    for fmt in ("txt", "md", "docx", "xlsx"):
        exported = export_markdown(markdown, export_format=fmt, source_name="scan.pdf")
        assert exported["format"] == fmt
        assert exported["file_name"].endswith(f".{fmt}")
        assert isinstance(exported["content"], (bytes, bytearray))
        assert len(exported["content"]) > 0


def test_signature_ink_transparent_crop():
    from PIL import Image

    from shared.doc_convert.render import RenderedPage
    from shared.doc_convert.structure import crop_signature_from_page

    image = Image.new("RGB", (200, 100), color=(255, 255, 255))
    for x in range(40, 160):
        for y in range(35, 55):
            image.putpixel((x, y), (10, 10, 10))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    page = RenderedPage(page_index=1, png_bytes=buffer.getvalue(), text_layer="", source_name="s.png")
    crop = crop_signature_from_page(
        page,
        bbox_norm=[0.15, 0.25, 0.85, 0.7],
        left_text="Руководитель",
        right_text="Иванов",
    )
    assert crop is not None
    with Image.open(io.BytesIO(crop.png_bytes)) as out:
        assert out.mode == "RGBA"
        # Corner should be transparent paper.
        assert out.getpixel((0, 0))[3] == 0


def test_guard_converts_listish_table():
    from shared.doc_convert.structure import guard_fix_structure

    blocks = [
        {
            "type": "table",
            "text": "",
            "align": "left",
            "bold": False,
            "indent_first": False,
            "items": [],
            "list_marker": "dash",
            "rows": [
                ["наименование квитанции;", "дата выдачи квитанции;"],
                ["стоимость такси;", "инн фрахтовщика;"],
                ["маршрут движения;", "класс тарифа такси."],
            ],
            "left_text": "",
            "right_text": "",
            "has_handwriting": False,
            "bbox_norm": [0.1, 0.4, 0.9, 0.7],
            "ink_bbox_norm": [0, 0, 0, 0],
        }
    ]
    fixed = guard_fix_structure(blocks)
    assert fixed[0]["type"] == "list"
    assert len(fixed[0]["items"]) == 6


def test_structure_to_markdown_and_docx_export():
    from shared.doc_convert.exporters import export_markdown
    from shared.doc_convert.models import SignatureCrop
    from shared.doc_convert.structure import structure_to_markdown

    payload = {
        "page_index": 1,
        "doc_kind": "form",
        "title": "Заявка на доступ",
        "title_centered": True,
        "title_bbox_norm": [0.2, 0.05, 0.8, 0.12],
        "blocks": [
            {
                "type": "table",
                "text": "",
                "align": "left",
                "bold": False,
                "indent_first": False,
                "items": [],
                "list_marker": "dash",
                "rows": [
                    ["Наименование подразделения", "Отдел учета"],
                    ["ФИО", "Сафронова Лилия Сергеевна"],
                ],
                "left_text": "",
                "right_text": "",
                "has_handwriting": False,
                "bbox_norm": [0.1, 0.2, 0.9, 0.55],
                "ink_bbox_norm": [0, 0, 0, 0],
            },
            {
                "type": "signature",
                "text": "",
                "align": "left",
                "bold": False,
                "indent_first": False,
                "items": [],
                "list_marker": "dash",
                "rows": [],
                "left_text": "Руководитель",
                "right_text": "Иванов И.И.",
                "has_handwriting": True,
                "bbox_norm": [0.1, 0.75, 0.9, 0.9],
                "ink_bbox_norm": [0.4, 0.8, 0.6, 0.9],
            },
        ],
    }
    markdown = structure_to_markdown(payload)
    assert "Заявка на доступ" in markdown
    assert "Сафронова Лилия Сергеевна" in markdown

    png = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f\x00"
        b"\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    exported = export_markdown(
        markdown,
        export_format="docx",
        source_name="form.pdf",
        structure_pages=[payload],
        signatures=[
            SignatureCrop(
                page_index=1,
                png_bytes=png,
                left_text="Руководитель",
                right_text="Иванов И.И.",
            )
        ],
    )
    assert exported["format"] == "docx"
    assert len(exported["content"]) > 1000


def test_structure_preview_and_verify_refine(monkeypatch):
    from shared.doc_convert import verify as verify_mod

    structure = {
        "page_index": 1,
        "doc_kind": "prose",
        "title": "Правила такси",
        "title_centered": True,
        "title_bbox_norm": [0.15, 0.05, 0.85, 0.12],
        "blocks": [
            {
                "type": "table",
                "text": "",
                "align": "left",
                "bold": False,
                "indent_first": False,
                "items": [],
                "list_marker": "dash",
                "rows": [["короткий ярлык", "Значение поля"]],
                "left_text": "",
                "right_text": "",
                "has_handwriting": False,
                "bbox_norm": [0.1, 0.2, 0.9, 0.4],
                "ink_bbox_norm": [0, 0, 0, 0],
            }
        ],
    }
    preview = verify_mod.render_structure_preview_png(structure)
    assert preview[:8] == b"\x89PNG\r\n\x1a\n"

    def fake_verify(**kwargs):
        return {
            "ok": False,
            "fidelity": 0.4,
            "issues": [{"code": "list_as_table", "detail": "list became table"}],
            "corrected": {
                "page_index": 1,
                "doc_kind": "prose",
                "title": "Правила такси",
                "title_centered": True,
                "title_bbox_norm": [0.15, 0.05, 0.85, 0.12],
                "blocks": [
                    {
                        "type": "list",
                        "text": "",
                        "align": "left",
                        "bold": False,
                        "indent_first": False,
                        "items": ["пункт 1", "пункт 2"],
                        "list_marker": "dash",
                        "rows": [],
                        "left_text": "",
                        "right_text": "",
                        "has_handwriting": False,
                        "bbox_norm": [0.1, 0.25, 0.9, 0.55],
                        "ink_bbox_norm": [0, 0, 0, 0],
                    }
                ],
            },
        }

    monkeypatch.setattr(verify_mod, "verify_structure_against_page", lambda **kwargs: fake_verify(**kwargs))
    monkeypatch.setattr(verify_mod, "VERIFY_ENABLED", True)
    monkeypatch.setattr(verify_mod, "VERIFY_ROUNDS", 1)
    monkeypatch.setattr(verify_mod, "REBUILD_ON_LOW_FIDELITY", False)
    monkeypatch.setattr(
        "shared.doc_convert.region_refine.enrich_structure_from_regions",
        lambda page_png, structure, **kwargs: (structure, []),
    )
    refined, warnings = verify_mod.refine_structure_with_verify(b"\x89PNG\r\n\x1a\n", structure)
    assert refined["blocks"][0]["type"] == "list"
    assert refined["blocks"][0]["items"] == ["пункт 1", "пункт 2"]
    assert warnings


def test_side_by_side_and_clamp_bloated():
    from PIL import Image

    from shared.doc_convert import verify as verify_mod

    left = Image.new("RGB", (200, 300), color=(255, 255, 255))
    right = Image.new("RGB", (200, 300), color=(240, 240, 255))
    left_buf = io.BytesIO()
    right_buf = io.BytesIO()
    left.save(left_buf, format="PNG")
    right.save(right_buf, format="PNG")
    combined = verify_mod._side_by_side_png(left_buf.getvalue(), right_buf.getvalue())
    assert combined[:8] == b"\x89PNG\r\n\x1a\n"
    with Image.open(io.BytesIO(combined)) as out:
        assert out.width > 200

    bloated = {
        "page_index": 1,
        "doc_kind": "form",
        "title": "X",
        "title_centered": True,
        "title_bbox_norm": [0, 0, 1, 0.1],
        "blocks": [
            {
                "type": "table",
                "text": "",
                "align": "left",
                "bold": False,
                "indent_first": False,
                "items": [],
                "list_marker": "dash",
                "rows": [["cell" * 200] * 20 for _ in range(80)],
                "left_text": "",
                "right_text": "",
                "has_handwriting": False,
                "bbox_norm": [0.1, 0.2, 0.9, 0.9],
                "ink_bbox_norm": [0, 0, 0, 0],
            }
        ],
    }
    trimmed = verify_mod.clamp_bloated_structure(bloated, max_chars=5000)
    from shared.doc_convert.structure import structure_to_markdown

    assert len(structure_to_markdown(trimmed)) < len(structure_to_markdown(bloated))


def test_prose_blocks_do_not_become_table():
    from shared.doc_convert.exporters import export_markdown
    from shared.doc_convert.structure import structure_to_markdown

    payload = {
        "page_index": 1,
        "doc_kind": "prose",
        "title": "Правила использования такси в служебной командировке.",
        "title_centered": True,
        "title_bbox_norm": [0.1, 0.04, 0.9, 0.1],
        "blocks": [
            {
                "type": "paragraph",
                "text": "Командируемый работник может использовать такси, __при невозможности воспользоваться общественным транспортом__, чтобы добраться.",
                "align": "justify",
                "bold": False,
                "indent_first": True,
                "items": [],
                "list_marker": "dash",
                "rows": [],
                "left_text": "",
                "right_text": "",
                "has_handwriting": False,
                "bbox_norm": [0.08, 0.12, 0.92, 0.28],
                "ink_bbox_norm": [0, 0, 0, 0],
            },
            {
                "type": "list",
                "text": "",
                "align": "left",
                "bold": False,
                "indent_first": False,
                "items": [
                    "наименование, серия и номер квитанции;",
                    "дата выдачи квитанции;",
                    "стоимость пользования легковым такси;",
                ],
                "list_marker": "dash",
                "rows": [],
                "left_text": "",
                "right_text": "",
                "has_handwriting": False,
                "bbox_norm": [0.1, 0.4, 0.9, 0.65],
                "ink_bbox_norm": [0, 0, 0, 0],
            },
            {
                "type": "paragraph",
                "text": "Нецелевое использование такси запрещено.",
                "align": "left",
                "bold": True,
                "indent_first": False,
                "items": [],
                "list_marker": "dash",
                "rows": [],
                "left_text": "",
                "right_text": "",
                "has_handwriting": False,
                "bbox_norm": [0.1, 0.8, 0.9, 0.9],
                "ink_bbox_norm": [0, 0, 0, 0],
            },
        ],
    }
    markdown = structure_to_markdown(payload)
    assert markdown.count("Правила использования такси") == 1
    assert "| " not in markdown
    assert "— наименование" in markdown or "- наименование" in markdown
    assert "__при невозможности" in markdown

    exported = export_markdown(
        markdown,
        export_format="docx",
        source_name="taxi.pdf",
        structure_pages=[payload],
    )
    assert exported["format"] == "docx"
    assert len(exported["content"]) > 1000


def test_export_xlsx_without_tables_warns():
    from openpyxl import load_workbook

    from shared.doc_convert.exporters import export_markdown

    exported = export_markdown("Just text\n\nSecond line", export_format="xlsx", source_name="note")
    assert exported["warnings"]
    assert exported["format"] == "xlsx"
    workbook = load_workbook(io.BytesIO(exported["content"]))
    assert workbook.sheetnames
    values = [cell.value for cell in workbook.active["A"] if cell.value]
    assert "Just text" in values or any("Just text" in str(v) for v in values)


def test_export_xlsx_from_structure_tables_and_text():
    from openpyxl import load_workbook

    from shared.doc_convert.exporters import export_markdown

    pages = [
        {
            "page_index": 1,
            "doc_kind": "form",
            "title": "Заявка на доступ",
            "title_centered": True,
            "title_bbox_norm": [0.1, 0.05, 0.9, 0.12],
            "blocks": [
                {
                    "type": "paragraph",
                    "text": "Прошу предоставить доступ к **1С**.",
                    "align": "left",
                    "bold": False,
                    "indent_first": False,
                    "items": [],
                    "list_marker": "dash",
                    "rows": [],
                    "left_text": "",
                    "right_text": "",
                    "has_handwriting": False,
                    "bbox_norm": [0.1, 0.15, 0.9, 0.22],
                    "ink_bbox_norm": [0, 0, 0, 0],
                },
                {
                    "type": "table",
                    "text": "",
                    "align": "left",
                    "bold": False,
                    "indent_first": False,
                    "items": [],
                    "list_marker": "dash",
                    "rows": [
                        ["ФИО", "Иванов И.И."],
                        ["Отдел", "IT"],
                        ["Срок", ""],
                    ],
                    "left_text": "",
                    "right_text": "",
                    "has_handwriting": False,
                    "bbox_norm": [0.1, 0.25, 0.9, 0.7],
                    "ink_bbox_norm": [0, 0, 0, 0],
                },
                {
                    "type": "list",
                    "text": "",
                    "align": "left",
                    "bold": False,
                    "indent_first": False,
                    "items": ["чтение", "запись"],
                    "list_marker": "dash",
                    "rows": [],
                    "left_text": "",
                    "right_text": "",
                    "has_handwriting": False,
                    "bbox_norm": [0.1, 0.72, 0.9, 0.85],
                    "ink_bbox_norm": [0, 0, 0, 0],
                },
            ],
        }
    ]
    exported = export_markdown(
        "# ignored markdown without tables",
        export_format="xlsx",
        source_name="access.pdf",
        structure_pages=pages,
    )
    assert exported["format"] == "xlsx"
    workbook = load_workbook(io.BytesIO(exported["content"]))
    assert "Документ" in workbook.sheetnames
    assert len(workbook.sheetnames) >= 2

    doc = workbook["Документ"]
    doc_values = " ".join(str(cell.value or "") for row in doc.iter_rows() for cell in row)
    assert "Заявка на доступ" in doc_values
    assert "1С" in doc_values
    assert "чтение" in doc_values

    table_sheet = next(name for name in workbook.sheetnames if name != "Документ")
    sheet = workbook[table_sheet]
    assert sheet["A1"].value == "Поле"
    assert sheet["B1"].value == "Значение"
    assert sheet["A2"].value == "ФИО"
    assert sheet["B2"].value == "Иванов И.И."
    assert sheet["A3"].value == "Отдел"


def test_export_xlsx_markdown_tables_styled():
    from openpyxl import load_workbook

    from shared.doc_convert.exporters import export_markdown

    markdown = (
        "# Report\n\nIntro text\n\n"
        "| Name | Qty |\n| --- | --- |\n| A | 1 |\n| B | 2 |\n"
    )
    exported = export_markdown(markdown, export_format="xlsx", source_name="report")
    workbook = load_workbook(io.BytesIO(exported["content"]))
    assert "Документ" in workbook.sheetnames
    assert any(name.startswith("Таблица") for name in workbook.sheetnames)
    table = workbook[[n for n in workbook.sheetnames if n.startswith("Таблица")][0]]
    assert table["A1"].value == "Name"
    assert table["B2"].value == "1"


def test_hybrid_text_page_skips_vision(monkeypatch):
    from shared.doc_convert.models import ConvertSource
    from shared.doc_convert.render import RenderedPage
    from shared.doc_convert import service as doc_service

    long_text = (
        "Это цифровой PDF с достаточным количеством связного текста на странице для hybrid-режима. "
        "Второй абзац тоже длинный и нормально читается без разбиения на отдельные слова."
    )

    def fake_render_source(**kwargs):
        return [
            RenderedPage(
                page_index=1,
                png_bytes=b"\x89PNG\r\n\x1a\n",
                text_layer=long_text,
                source_name="digital.pdf",
            )
        ]

    called = {"vision": 0}

    def fake_vision(**kwargs):
        called["vision"] += 1
        return ("SHOULD_NOT_USE", {})

    monkeypatch.setattr(doc_service, "render_source", fake_render_source)
    result = doc_service.convert_sources(
        [ConvertSource(file_name="digital.pdf", data=b"%PDF-1.4", mime_type="application/pdf")],
        vision_fn=fake_vision,
    )
    assert called["vision"] == 0
    assert "цифровой PDF" in result.markdown
    assert result.used_vision is False


def test_fragmented_pdf_text_falls_back_to_vision(monkeypatch):
    from shared.doc_convert.models import ConvertSource
    from shared.doc_convert.render import RenderedPage, page_has_usable_text
    from shared.doc_convert import service as doc_service

    fragmented = "\n".join(
        [
            "Наименование",
            "подразделения",
            "Отдел",
            "учета",
            "движения",
            "ТМЦ",
            "Должность",
            "работника",
            "Кладовщик",
            "ФИО",
            "Сафронова",
            "Лилия",
            "Сергеевна",
            "Обоснование",
            "доступа",
            "к",
            "1С",
        ]
    )
    assert page_has_usable_text(fragmented) is False

    def fake_render_source(**kwargs):
        return [
            RenderedPage(
                page_index=1,
                png_bytes=b"\x89PNG\r\n\x1a\n",
                text_layer=fragmented,
                source_name="form.pdf",
            )
        ]

    def fake_vision(**kwargs):
        return (
            "# Заявка\n\n| Поле | Значение |\n| --- | --- |\n| ФИО | Сафронова Лилия Сергеевна |\n",
            {},
        )

    monkeypatch.setattr(doc_service, "render_source", fake_render_source)
    result = doc_service.convert_sources(
        [ConvertSource(file_name="form.pdf", data=b"%PDF-1.4", mime_type="application/pdf")],
        vision_fn=fake_vision,
    )
    assert result.used_vision is True
    assert "Сафронова Лилия Сергеевна" in result.markdown
    assert "| ФИО |" in result.markdown


def test_image_source_uses_vision(monkeypatch):
    from shared.doc_convert.models import ConvertSource
    from shared.doc_convert.render import RenderedPage
    from shared.doc_convert import service as doc_service

    def fake_render_source(**kwargs):
        return [
            RenderedPage(
                page_index=1,
                png_bytes=b"\x89PNG\r\n\x1a\n",
                text_layer="",
                source_name="scan.png",
            )
        ]

    def fake_vision(**kwargs):
        return ("# Scanned\n\nBody line", {})

    monkeypatch.setattr(doc_service, "render_source", fake_render_source)
    # Pretend image bytes are valid by skipping real image decode via fake render.
    result = doc_service.convert_sources(
        [ConvertSource(file_name="scan.png", data=b"fake-image", mime_type="image/png")],
        vision_fn=fake_vision,
        force_vision=True,
    )
    assert "# Scanned" in result.markdown
    assert result.used_vision is True


def test_detect_format_from_text():
    from backend.ai_chat.doc_convert_runtime import detect_format_from_text

    assert detect_format_from_text("сделай в word") == "docx"
    assert detect_format_from_text("нужен excel") == "xlsx"
    assert detect_format_from_text("markdown please") == "md"
    assert detect_format_from_text("просто файл") is None


def test_doc_convert_format_choice_payload():
    from backend.ai_chat.action_cards import (
        ACTION_DOC_CONVERT_FORMAT_CHOICE,
        _build_doc_convert_format_preview,
        _normalize_doc_convert_format,
        _normalize_doc_convert_payload,
    )

    payload = _normalize_doc_convert_payload(
        {
            "markdown": "# Hello",
            "title": "Scan",
            "page_count": 2,
            "table_count": 1,
            "used_vision": True,
            "formats": ["docx", "xlsx"],
        }
    )
    preview = _build_doc_convert_format_preview(payload)
    assert preview["title"]
    assert preview["formats"] == ["docx", "xlsx"]
    assert _normalize_doc_convert_format("word") == "docx"
    assert ACTION_DOC_CONVERT_FORMAT_CHOICE == "ai.doc.convert.format_choice"


def test_convert_tool_needs_format_choice(monkeypatch):
    from backend.ai_chat.tools.context import AI_TOOL_FILES_CONVERT_DOCUMENT, AiToolExecutionContext
    from backend.ai_chat.tools.files import FilesConvertDocumentArgs, FilesConvertDocumentTool

    monkeypatch.setattr(
        "backend.ai_chat.doc_convert_runtime.convert_attachments_to_markdown",
        lambda **kwargs: {
            "markdown": "# Doc\n\nText",
            "source_names": ["a.pdf"],
            "warnings": [],
            "table_count": 0,
            "page_count": 1,
            "used_vision": True,
            "formats": ["docx", "txt", "md", "pdf", "xlsx"],
            "title": "a",
        },
    )
    tool = FilesConvertDocumentTool()
    context = AiToolExecutionContext(
        bot_id="b",
        bot_title="Converter",
        conversation_id="c",
        run_id="r",
        user_id=1,
        user_payload={"id": 1, "role": "viewer"},
        effective_database_id=None,
        enabled_tools=[AI_TOOL_FILES_CONVERT_DOCUMENT],
        tool_settings={},
        allow_generated_artifacts=True,
        trigger_message_id="msg-1",
    )
    result = tool.execute(context=context, args=FilesConvertDocumentArgs())
    assert result.ok is True
    assert result.data["needs_format_choice"] is True
    assert result.data["markdown"].startswith("# Doc")
