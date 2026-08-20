from __future__ import annotations

import importlib
import sys
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest

PROJECT_ROOT = Path(__file__).resolve().parents[1]
WEB_ROOT = PROJECT_ROOT / "WEB-itinvent"

if str(WEB_ROOT) not in sys.path:
    sys.path.insert(0, str(WEB_ROOT))

try:
    import openpyxl
except Exception:  # pragma: no cover
    openpyxl = None


pytestmark = pytest.mark.skipif(openpyxl is None, reason="openpyxl is required")


def _workbook_bytes(sheets: list[dict]) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for sheet in sheets:
        worksheet = workbook.create_sheet(title=sheet["title"])
        header_row = int(sheet.get("header_row") or 1)
        headers = sheet["headers"]
        for column, header in enumerate(headers, start=1):
            worksheet.cell(header_row, column).value = header
        for offset, row in enumerate(sheet.get("rows") or [], start=1):
            for column, value in enumerate(row, start=1):
                worksheet.cell(header_row + offset, column).value = value
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _row_dict(row) -> dict:
    if hasattr(row, "keys"):
        return {str(key): row[key] for key in row.keys()}
    return dict(row)


@pytest.fixture
def isolated_network_service(temp_dir, monkeypatch):
    network_service_module = importlib.import_module("backend.services.network_service")
    monkeypatch.setattr(network_service_module, "is_app_database_configured", lambda: False)
    monkeypatch.setattr(
        network_service_module,
        "get_local_store",
        lambda: SimpleNamespace(db_path=str(Path(temp_dir) / "network_equipment_import.db")),
    )
    service = network_service_module.NetworkService()
    branch = service.ensure_branch(
        city_code="tmn",
        branch_code="tmn-velizh-test",
        name="Велижанский тракт 6 км.",
    )
    return service, int(branch["id"])


def _device_rows(service, branch_id: int) -> list[dict]:
    with service._lock, service._connect() as conn:
        rows = conn.execute(
            """
            SELECT device_code, sheet_name, model,
                   (SELECT COUNT(*) FROM network_ports p WHERE p.device_id = network_devices.id) AS ports
            FROM network_devices
            WHERE branch_id=?
            ORDER BY id
            """,
            (branch_id,),
        ).fetchall()
        return [_row_dict(row) for row in rows]


def test_equipment_import_keeps_colliding_switch_sheets_separate(isolated_network_service):
    service, branch_id = isolated_network_service
    payload = _workbook_bytes(
        [
            {
                "title": "eltex mes 2348b (1)",
                "headers": ["Swich", "Port", "Name"],
                "rows": [["2348b", "Gi1/0/1", "pc-a"], ["2348b", "Gi1/0/2", "pc-b"]],
            },
            {
                "title": "eltex mes 2348b (3)",
                "headers": ["Swich", "Port", "Name"],
                "rows": [["2348b", "Gi1/0/1", "pc-c"], ["2348b", "Gi1/0/2", "pc-d"]],
            },
            {
                "title": "cisco офис 1 этаж",
                "headers": ["Switch", "Port", "Name"],
                "rows": [["cisco офис 1 этаж", "Gi0/1", "phone"]],
            },
        ]
    )

    result = service.import_equipment_from_excel(
        branch_id=branch_id,
        file_name="схема сети.xlsx",
        file_bytes=payload,
        actor_user_id=1,
        actor_role="admin",
    )

    summary = result["summary"]
    assert summary["sheets_total"] == 3
    assert summary["sheets_imported"] == 3
    assert summary["devices_created"] == 3
    assert summary["ports_created"] == 5

    devices = _device_rows(service, branch_id)
    codes = {item["device_code"] for item in devices}
    assert codes == {"eltex mes 2348b (1)", "eltex mes 2348b (3)", "cisco офис 1 этаж"}
    by_code = {item["device_code"]: item for item in devices}
    assert by_code["eltex mes 2348b (1)"]["ports"] == 2
    assert by_code["eltex mes 2348b (3)"]["ports"] == 2
    assert by_code["eltex mes 2348b (1)"]["model"] == "2348b"
    assert by_code["cisco офис 1 этаж"]["ports"] == 1


def test_equipment_import_skips_sheets_without_port_column(isolated_network_service):
    service, branch_id = isolated_network_service
    payload = _workbook_bytes(
        [
            {
                "title": "notes",
                "headers": ["Comment"],
                "rows": [["ignore me"]],
            },
            {
                "title": "at-8000s",
                "headers": ["ASW", "Port"],
                "rows": [["at-8000s", "1"], ["at-8000s", "2"]],
            },
        ]
    )

    result = service.import_equipment_from_excel(
        branch_id=branch_id,
        file_name="схема сети.xlsx",
        file_bytes=payload,
        actor_user_id=None,
        actor_role=None,
    )

    summary = result["summary"]
    assert summary["sheets_total"] == 2
    assert summary["sheets_imported"] == 1
    assert summary["sheets_skipped"] == 1
    assert summary["skipped_sheets"] == ["notes"]
    devices = _device_rows(service, branch_id)
    assert [item["device_code"] for item in devices] == ["at-8000s"]
    assert devices[0]["ports"] == 2


def test_equipment_import_reads_header_below_title_row(isolated_network_service):
    service, branch_id = isolated_network_service
    payload = _workbook_bytes(
        [
            {
                "title": "cas125-24g рмм",
                "header_row": 2,
                "headers": ["Swich", "Port", "Name"],
                "rows": [["2324b", "1", "host-1"]],
            }
        ]
    )

    result = service.import_equipment_from_excel(
        branch_id=branch_id,
        file_name="схема сети.xlsx",
        file_bytes=payload,
        actor_user_id=None,
        actor_role=None,
    )

    assert result["summary"]["ports_created"] == 1
    devices = _device_rows(service, branch_id)
    assert devices[0]["device_code"] == "2324b"
    assert devices[0]["sheet_name"] == "cas125-24g рмм"
    assert devices[0]["ports"] == 1


def test_equipment_import_splits_sequential_switch_rows_on_one_sheet(isolated_network_service):
    service, branch_id = isolated_network_service
    payload = _workbook_bytes(
        [
            {
                "title": "access",
                "headers": ["Swich", "Port", "Port P/P", "Name"],
                "rows": [
                    ["Aruba 1", "1", "A-1", "ap-1"],
                    [None, "2", "A-2", "ap-1b"],
                    ["Aruba 2", "1", "B-1", "ap-2"],
                    ["Aruba 3", "1", "C-1", "ap-3"],
                ],
            }
        ]
    )

    result = service.import_equipment_from_excel(
        branch_id=branch_id,
        file_name="схема сети.xlsx",
        file_bytes=payload,
        actor_user_id=1,
        actor_role="admin",
    )

    summary = result["summary"]
    assert summary["sheets_imported"] == 1
    assert summary["devices_created"] == 3
    assert summary["ports_created"] == 4

    devices = _device_rows(service, branch_id)
    by_code = {item["device_code"]: item for item in devices}
    assert set(by_code) == {"Aruba 1", "Aruba 2", "Aruba 3"}
    assert by_code["Aruba 1"]["ports"] == 2
    assert by_code["Aruba 2"]["ports"] == 1
    assert by_code["Aruba 3"]["ports"] == 1
    assert by_code["Aruba 1"]["sheet_name"] == "access"

    with service._lock, service._connect() as conn:
        sockets = conn.execute(
            """
            SELECT s.socket_code, d.device_code
            FROM network_sockets s
            JOIN network_devices d ON d.id=s.device_id
            WHERE s.branch_id=?
            ORDER BY s.socket_code
            """,
            (branch_id,),
        ).fetchall()
        pairs = {(_row_dict(row)["socket_code"], _row_dict(row)["device_code"]) for row in sockets}
        assert ("A-1", "Aruba 1") in pairs
        assert ("B-1", "Aruba 2") in pairs
        assert ("C-1", "Aruba 3") in pairs


def test_socket_sync_does_not_steal_socket_from_another_device(isolated_network_service):
    service, branch_id = isolated_network_service
    first = service.create_device(
        branch_id=branch_id,
        payload={"device_code": "Aruba 1", "device_type": "switch"},
        actor_user_id=1,
        actor_role="admin",
    )
    second = service.create_device(
        branch_id=branch_id,
        payload={"device_code": "Aruba 2", "device_type": "switch"},
        actor_user_id=1,
        actor_role="admin",
    )
    service.create_port(
        device_id=int(first["id"]),
        payload={"port_name": "1", "patch_panel_port": "1-1"},
        actor_user_id=1,
        actor_role="admin",
    )
    service.create_port(
        device_id=int(second["id"]),
        payload={"port_name": "1", "patch_panel_port": "1-1"},
        actor_user_id=1,
        actor_role="admin",
    )

    with service._lock, service._connect() as conn:
        sockets = conn.execute(
            """
            SELECT s.socket_code, s.device_id, d.device_code, p.port_name
            FROM network_sockets s
            JOIN network_devices d ON d.id=s.device_id
            LEFT JOIN network_ports p ON p.id=s.port_id
            WHERE s.branch_id=? AND s.socket_code=?
            """,
            (branch_id, "1-1"),
        ).fetchall()
        assert len(sockets) == 1
        owner = _row_dict(sockets[0])
        assert owner["device_code"] == "Aruba 1"

        second_port = conn.execute(
            """
            SELECT p.id, s.id socket_id
            FROM network_ports p
            LEFT JOIN network_sockets s ON s.port_id=p.id
            WHERE p.device_id=?
            """,
            (int(second["id"]),),
        ).fetchone()
        second_port = _row_dict(second_port)
        assert second_port["socket_id"] is None


def test_sockets_import_links_sequential_aruba_rows(isolated_network_service):
    service, branch_id = isolated_network_service
    equipment = _workbook_bytes(
        [
            {
                "title": "access",
                "headers": ["Swich", "Port", "Port P/P"],
                "rows": [
                    ["Aruba 1", "1", "ROOM-1"],
                    ["Aruba 2", "1", "ROOM-2"],
                ],
            }
        ]
    )
    service.import_equipment_from_excel(
        branch_id=branch_id,
        file_name="equipment.xlsx",
        file_bytes=equipment,
        actor_user_id=1,
        actor_role="admin",
    )
    sockets = _workbook_bytes(
        [
            {
                "title": "розетки",
                "headers": ["ASW", "Port", "Socket", "FIO"],
                "rows": [
                    ["Aruba 1", "1", "ROOM-1", "Иванов"],
                    [None, "1", "ROOM-1", "Иванов"],
                    ["Aruba 2", "1", "ROOM-2", "Петров"],
                ],
            }
        ]
    )
    result = service.import_sockets_template(
        branch_id=branch_id,
        file_name="sockets.xlsx",
        file_bytes=sockets,
        actor_user_id=1,
        actor_role="admin",
    )
    assert result["updated"] >= 2

    with service._lock, service._connect() as conn:
        rows = conn.execute(
            """
            SELECT s.socket_code, d.device_code, s.fio
            FROM network_sockets s
            JOIN network_devices d ON d.id=s.device_id
            WHERE s.branch_id=?
            ORDER BY s.socket_code
            """,
            (branch_id,),
        ).fetchall()
        by_code = {_row_dict(row)["socket_code"]: _row_dict(row) for row in rows}
        assert by_code["ROOM-1"]["device_code"] == "Aruba 1"
        assert by_code["ROOM-2"]["device_code"] == "Aruba 2"
