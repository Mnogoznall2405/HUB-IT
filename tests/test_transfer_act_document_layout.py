from __future__ import annotations

import importlib
import shutil
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn


PROJECT_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = PROJECT_ROOT / "templates" / "docx_transfer_act.docx"
EXPECTED_HEADERS = [
    "ID",
    "Тип оборудования",
    "Модель оборудования",
    "Серийный номер",
    "Номер партий",
    "Инвентарный номер",
]
EXPECTED_GRID_WIDTHS = [568, 2100, 2100, 2500, 1600, 2178]
TO_EMPLOYEE = "Козловский Анатолий Максимович"
FROM_EMPLOYEE = "Козловский Максим Евгеньевич"
DEPARTMENT = "Отдел по защите информации"


def _assert_transfer_act_layout(path: Path) -> None:
    document = Document(path)

    equipment_table = document.tables[0]
    assert len(equipment_table.columns) == 6
    assert [cell.text for cell in equipment_table.rows[0].cells] == EXPECTED_HEADERS
    assert [cell.text for cell in equipment_table.rows[1].cells] == [
        "1",
        "Системный блок",
        "Lenovo ThinkCentre M720q",
        "PC1L9HLE",
        "ЦБ-00087978",
        "101646",
    ]
    grid_widths = [
        int(grid_column.get(qn("w:w")))
        for grid_column in equipment_table._tbl.tblGrid.gridCol_lst
    ]
    assert grid_widths == EXPECTED_GRID_WIDTHS

    document_text = "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )
    document_text += "\n" + "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    assert DEPARTMENT not in document_text
    assert "{{TO_EMPLOYEE}}" not in document_text
    assert "{{FROM_EMPLOYEE}}" not in document_text

    signature_container = document.tables[1]
    assert len(signature_container.rows) == 1
    assert len(signature_container.columns) == 1
    assert signature_container.rows[0]._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None

    signature_grid = signature_container.cell(0, 0).tables[0]
    assert len(signature_grid.rows) == 4
    assert len(signature_grid.columns) == 3
    assert [cell.text for cell in signature_grid.rows[0].cells] == [
        "С условиями ознакомлен:",
        "",
        "Выдал:",
    ]
    assert [cell.text for cell in signature_grid.rows[1].cells] == [
        TO_EMPLOYEE,
        "",
        FROM_EMPLOYEE,
    ]
    assert [cell.text for cell in signature_grid.rows[3].cells] == [
        "(подпись)",
        "",
        "(подпись)",
    ]


def test_web_transfer_act_has_six_columns_and_atomic_signature_block(tmp_path, monkeypatch):
    transfer_service = importlib.import_module("backend.services.transfer_service")
    monkeypatch.setattr(transfer_service, "DEFAULT_ACTS_DIR", tmp_path)

    path, file_type = transfer_service._build_docx_act(
        old_employee=FROM_EMPLOYEE,
        new_employee=TO_EMPLOYEE,
        new_employee_dept=DEPARTMENT,
        items=[
            {
                "type_name": "Системный блок",
                "model_name": "Lenovo ThinkCentre M720q",
                "serial_no": "PC1L9HLE",
                "part_no": "ЦБ-00087978",
                "inv_no": "101646",
            }
        ],
        deterministic_act_id="layout-regression-web",
    )

    assert file_type == "docx"
    _assert_transfer_act_layout(path)


@pytest.mark.asyncio
async def test_bot_transfer_act_has_six_columns_and_atomic_signature_block(tmp_path, monkeypatch):
    pdf_generator = importlib.import_module("bot.services.pdf_generator")
    template_dir = tmp_path / "templates"
    template_dir.mkdir()
    shutil.copy2(TEMPLATE_PATH, template_dir / TEMPLATE_PATH.name)
    monkeypatch.chdir(tmp_path)

    import docx2pdf

    def fail_conversion(*_args, **_kwargs):
        raise RuntimeError("conversion intentionally disabled for DOCX structure test")

    monkeypatch.setattr(docx2pdf, "convert", fail_conversion)

    generated_path = await pdf_generator.generate_transfer_act_pdf(
        new_employee=TO_EMPLOYEE,
        new_employee_dept=DEPARTMENT,
        old_employee=FROM_EMPLOYEE,
        serials_data=[
            {
                "serial": "PC1L9HLE",
                "equipment": {
                    "TYPE_NAME": "Системный блок",
                    "MODEL_NAME": "Lenovo ThinkCentre M720q",
                    "PART_NO": "ЦБ-00087978",
                    "INV_NO": "101646",
                },
            }
        ],
        db_name="main",
    )

    assert generated_path is not None
    assert Path(generated_path).suffix == ".docx"
    _assert_transfer_act_layout(Path(generated_path))
