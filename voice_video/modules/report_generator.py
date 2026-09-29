import logging
import json
import re
from pathlib import Path
from typing import Dict, List, Any, Optional
from datetime import datetime
import html
from xml.sax.saxutils import escape as xml_escape

from .audio_processor import ProcessingConfig
from config import OUTPUT_DIR

logger = logging.getLogger(__name__)

class ReportGenerator:
    """Класс для генерации отчетов в различных форматах"""
    
    def __init__(self, config: ProcessingConfig):
        self.config = config
        self.output_dir = Path(OUTPUT_DIR)
        self.output_dir.mkdir(exist_ok=True)
    
    def generate_all_reports(self,
                           analysis_results: Dict[str, Any],
                           segments_with_speakers: List[Dict],
                           insights: Dict[str, Any],
                           base_filename: str,
                           word_segments: Optional[List[Dict]] = None) -> Dict[str, str]:
        """Генерирует отчеты во всех форматах."""
        try:
            logger.info(f"📄 Генерация отчетов для {base_filename}...")
            
            generated_files = {}

            # Все отчёты совещания — в отдельной папке output/<base>/
            root_output_dir = self.output_dir
            meeting_dir = root_output_dir / base_filename
            meeting_dir.mkdir(parents=True, exist_ok=True)
            self.output_dir = meeting_dir
            try:
                # Подготавливаем данные для отчетов
                report_data = self._prepare_report_data(analysis_results, segments_with_speakers, insights)

                # Генерируем JSON отчет
                json_path = self._generate_json_report(report_data, base_filename)
                if json_path:
                    generated_files['json'] = json_path

                # Генерируем Markdown отчет
                md_path = self._generate_markdown_report(report_data, base_filename)
                if md_path:
                    generated_files['markdown'] = md_path

                # Генерируем HTML отчет
                html_path = self._generate_html_report(report_data, base_filename)
                if html_path:
                    generated_files['html'] = html_path

                # Генерируем PDF отчет (если возможно)
                pdf_path = self._generate_pdf_report(report_data, html_path, base_filename)
                if pdf_path:
                    generated_files['pdf'] = pdf_path

                # Генерируем DOCX отчеты (полный анализ + отдельный протокол)
                docx_files = self._generate_docx_reports(report_data, base_filename)
                generated_files.update(docx_files)

                # Генерируем отдельный файл транскрипции
                transcript_path = self._generate_transcript_file(segments_with_speakers, base_filename, word_segments)
                if transcript_path:
                    generated_files['transcript'] = transcript_path
            finally:
                self.output_dir = root_output_dir

            logger.info(f"✅ Сгенерировано {len(generated_files)} файлов отчетов в {meeting_dir}")
            return generated_files
            
        except Exception as e:
            logger.error(f"❌ Ошибка генерации отчетов: {e}")
            return {}
    
    def _prepare_report_data(self, 
                           analysis_results: Dict[str, Any], 
                           segments_with_speakers: List[Dict],
                           insights: Dict[str, Any]) -> Dict[str, Any]:
        """Подготавливает данные для генерации отчетов."""
        try:
            # Базовая информация
            metadata = analysis_results.get('metadata', {})
            statistics = analysis_results.get('statistics', {})
            
            # Форматируем участников
            participants = self._format_participants(statistics.get('speaker_statistics', {}))
            
            # Подготавливаем основные разделы
            report_data = {
                'metadata': {
                    'title': 'Анализ встречи',
                    'generated_at': datetime.now().isoformat(),
                    'analysis_timestamp': metadata.get('analysis_timestamp', ''),
                    'total_duration': statistics.get('meeting_duration_formatted', ''),
                    'total_segments': metadata.get('total_segments', 0),
                    'unique_speakers': statistics.get('unique_speakers', 0)
                },
                'participants': participants,
                'protocol': (analysis_results.get('detailed_analysis') or {}).get('protocol_markdown', ''),
                'action_registry': (analysis_results.get('detailed_analysis') or {}).get('action_registry', ''),
                'topics': (analysis_results.get('detailed_analysis') or {}).get('topics', []),
                'summary': analysis_results.get('summary', ''),
                'action_items': analysis_results.get('action_items', ''),
                'questions': analysis_results.get('questions', ''),
                'participants_analysis': analysis_results.get('participants', ''),
                'statistics': statistics,
                'insights': insights,
                'transcript': self._format_transcript(segments_with_speakers),
                'keywords': insights.get('keywords', [])[:10]  # Топ 10 ключевых слов
            }
            
            return report_data
            
        except Exception as e:
            logger.error(f"❌ Ошибка подготовки данных отчета: {e}")
            return {}
    
    def _format_participants(self, speaker_stats: Dict[str, Any]) -> List[Dict[str, Any]]:
        """Форматирует информацию об участниках."""
        try:
            participants = []
            
            for speaker, stats in speaker_stats.items():
                participant = {
                    'name': speaker,
                    'segments_count': stats.get('segments_count', 0),
                    'total_duration_seconds': stats.get('total_duration', 0),
                    'participation_percentage': round(stats.get('participation_percentage', 0), 1),
                    'words_count': stats.get('words_count', 0),
                    'words_per_minute': round(stats.get('words_per_minute', 0), 1),
                    'avg_segment_duration': round(stats.get('avg_segment_duration', 0), 1)
                }
                participants.append(participant)
            
            # Сортируем по проценту участия
            participants.sort(key=lambda x: x['participation_percentage'], reverse=True)
            
            return participants
            
        except Exception as e:
            logger.error(f"❌ Ошибка форматирования участников: {e}")
            return []
    
    def _format_transcript(self, segments_with_speakers: List[Dict], format_type: str = 'detailed') -> List[Dict[str, Any]]:
        """Форматирует транскрипцию для отчета."""
        try:
            formatted_transcript = []
            
            for segment in segments_with_speakers:
                speaker = segment.get('speaker', 'UNKNOWN')
                text = segment.get('text', '').strip()
                start_time = segment.get('start', 0)
                end_time = segment.get('end', 0)
                
                if text:
                    formatted_segment = {
                        'speaker': speaker,
                        'text': text,
                        'start_time': start_time,
                        'end_time': end_time,
                        'start_time_formatted': self._format_time(start_time),
                        'end_time_formatted': self._format_time(end_time),
                        'duration': round(end_time - start_time, 2)
                    }
                    
                    formatted_transcript.append(formatted_segment)
            
            return formatted_transcript
            
        except Exception as e:
            logger.error(f"❌ Ошибка форматирования транскрипции: {e}")
            return []
    
    def _format_time(self, seconds: float) -> str:
        """Форматирует время в читаемый вид."""
        try:
            hours = int(seconds // 3600)
            minutes = int((seconds % 3600) // 60)
            secs = int(seconds % 60)
            
            if hours > 0:
                return f"{hours:02d}:{minutes:02d}:{secs:02d}"
            else:
                return f"{minutes:02d}:{secs:02d}"
                
        except Exception:
            return "00:00"
    
    def _generate_json_report(self, report_data: Dict[str, Any], base_filename: str) -> Optional[str]:
        """Генерирует JSON отчет."""
        try:
            output_path = self.output_dir / f"{base_filename}_report.json"
            
            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(report_data, f, ensure_ascii=False, indent=2)
            
            logger.info(f"✅ JSON отчет сохранен: {output_path}")
            return str(output_path)
            
        except Exception as e:
            logger.error(f"❌ Ошибка генерации JSON отчета: {e}")
            return None
    
    def _generate_markdown_report(self, report_data: Dict[str, Any], base_filename: str) -> Optional[str]:
        """Генерирует Markdown отчет."""
        try:
            output_path = self.output_dir / f"{base_filename}_report.md"
            
            md_content = self._build_markdown_content(report_data)
            
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(md_content)
            
            logger.info(f"✅ Markdown отчет сохранен: {output_path}")
            return str(output_path)
            
        except Exception as e:
            logger.error(f"❌ Ошибка генерации Markdown отчета: {e}")
            return None
    
    def _build_markdown_content(self, report_data: Dict[str, Any]) -> str:
        """Строит содержимое Markdown отчета."""
        try:
            metadata = report_data.get('metadata', {})
            participants = report_data.get('participants', [])
            transcript = report_data.get('transcript', [])
            keywords = report_data.get('keywords', [])
            insights = report_data.get('insights', {})
            
            md_lines = []

            # Заголовок
            md_lines.append(f"# {metadata.get('title', 'Анализ встречи')}")
            md_lines.append("")
            
            # Метаданные
            md_lines.append("## Информация о встрече")
            md_lines.append("")
            md_lines.append(f"- **Дата анализа:** {metadata.get('generated_at', '')}")
            md_lines.append(f"- **Длительность:** {metadata.get('total_duration', '')}")
            md_lines.append(f"- **Количество участников:** {metadata.get('unique_speakers', 0)}")
            md_lines.append(f"- **Всего сегментов:** {metadata.get('total_segments', 0)}")
            md_lines.append("")
            
            # Участники
            if participants:
                md_lines.append("## Участники")
                md_lines.append("")
                for participant in participants:
                    name = participant.get('name', 'Неизвестный')
                    md_lines.append(f"- {name}")
                md_lines.append("")
            
            # Реестр поручений — первый раздел отчёта
            action_registry = report_data.get('action_registry', '')
            if action_registry and action_registry.strip():
                registry_text = re.sub(r'(?m)^(#{1,5})\s', r'#\1 ', action_registry.strip())
                md_lines.append(registry_text)
                md_lines.append("")

            # Протокол совещания
            protocol = report_data.get('protocol', '')
            if protocol and protocol.strip():
                # Понижаем уровни заголовков протокола, чтобы он встроился в отчёт
                embedded = re.sub(r'(?m)^(#{1,5})\s', r'#\1 ', protocol.strip())
                md_lines.append(embedded)
                md_lines.append("")

            # Резюме
            summary = report_data.get('summary', '')
            if summary and summary.strip():
                md_lines.append("## Резюме встречи")
                md_lines.append("")
                md_lines.append(summary)
                md_lines.append("")
            
            # Открытые вопросы
            questions = report_data.get('questions', '')
            if questions and questions.strip():
                md_lines.append("## Открытые вопросы")
                md_lines.append("")
                md_lines.append(questions)
                md_lines.append("")

            participants_analysis = report_data.get('participants_analysis', '')
            if participants_analysis and participants_analysis.strip():
                md_lines.append("## Анализ участников")
                md_lines.append("")
                md_lines.append(participants_analysis)
                md_lines.append("")
            
            # Инсайты
            if insights:
                md_lines.append("## Анализ и инсайты")
                md_lines.append("")
                
                # Динамика разговора
                dynamics = insights.get('conversation_dynamics', {})
                if dynamics:
                    md_lines.append("### Динамика разговора")
                    md_lines.append("")
                    balance = dynamics.get('conversation_balance', '')
                    interactivity = dynamics.get('interactivity_level', '')
                    switches = dynamics.get('total_speaker_switches', 0)
                    
                    md_lines.append(f"- **Баланс разговора:** {balance}")
                    md_lines.append(f"- **Уровень интерактивности:** {interactivity}")
                    md_lines.append(f"- **Переключения между спикерами:** {switches}")
                    md_lines.append("")
                
                # Рекомендации
                recommendations = insights.get('recommendations', [])
                if recommendations:
                    md_lines.append("### Рекомендации")
                    md_lines.append("")
                    for rec in recommendations:
                        md_lines.append(f"- {rec}")
                    md_lines.append("")
            
            # Полная транскрипция
            if transcript:
                md_lines.append("## Полная транскрипция")
                md_lines.append("")
                
                for segment in transcript:
                    speaker = segment.get('speaker', 'UNKNOWN')
                    text = segment.get('text', '')
                    time_formatted = segment.get('start_time_formatted', '')
                    
                    md_lines.append(f"**[{time_formatted}] {speaker}:** {text}")
                    md_lines.append("")
            
            return "\n".join(md_lines)
            
        except Exception as e:
            logger.error(f"❌ Ошибка построения Markdown содержимого: {e}")
            return f"# Ошибка генерации отчета\n\nПроизошла ошибка: {str(e)}"
    
    def _generate_html_report(self, report_data: Dict[str, Any], base_filename: str) -> Optional[str]:
        """Генерирует HTML отчет."""
        try:
            output_path = self.output_dir / f"{base_filename}_report.html"
            
            html_content = self._build_html_content(report_data)
            
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(html_content)
            
            logger.info(f"✅ HTML отчет сохранен: {output_path}")
            return str(output_path)
            
        except Exception as e:
            logger.error(f"❌ Ошибка генерации HTML отчета: {e}")
            return None
    
    def _build_html_content(self, report_data: Dict[str, Any]) -> str:
        """Строит содержимое HTML отчета."""
        try:
            metadata = report_data.get('metadata', {})
            participants = report_data.get('participants', [])
            transcript = report_data.get('transcript', [])
            keywords = report_data.get('keywords', [])
            insights = report_data.get('insights', {})
            
            # Начало HTML документа
            html_parts = []
            html_parts.append('<!DOCTYPE html>')
            html_parts.append('<html lang="ru">')
            html_parts.append('<head>')
            html_parts.append('<meta charset="UTF-8">')
            html_parts.append('<meta name="viewport" content="width=device-width, initial-scale=1.0">')
            html_parts.append(f'<title>{html.escape(metadata.get("title", "Анализ встречи"))}</title>')
            
            # CSS стили
            html_parts.append('<style>')
            html_parts.append(self._get_html_styles())
            html_parts.append('</style>')
            html_parts.append('</head>')
            html_parts.append('<body>')
            
            # Контейнер
            html_parts.append('<div class="container">')
            
            # Заголовок
            html_parts.append(f'<h1>{html.escape(metadata.get("title", "Анализ встречи"))}</h1>')
            
            # Метаданные
            html_parts.append('<div class="metadata">')
            html_parts.append('<h2>Информация о встрече</h2>')
            html_parts.append('<div class="info-grid">')
            html_parts.append(f'<div class="info-item"><strong>Дата анализа:</strong> {html.escape(metadata.get("generated_at", ""))}</div>')
            html_parts.append(f'<div class="info-item"><strong>Длительность:</strong> {html.escape(metadata.get("total_duration", ""))}</div>')
            html_parts.append(f'<div class="info-item"><strong>Участники:</strong> {metadata.get("unique_speakers", 0)}</div>')
            html_parts.append(f'<div class="info-item"><strong>Сегменты:</strong> {metadata.get("total_segments", 0)}</div>')
            html_parts.append('</div>')
            html_parts.append('</div>')
            
            # Участники
            if participants:
                html_parts.append('<div class="participants">')
                html_parts.append('<h2>Участники</h2>')
                html_parts.append('<ul class="participants-list">')
                for participant in participants:
                    name = html.escape(participant.get('name', 'Неизвестный'))
                    html_parts.append(f'<li>{name}</li>')
                html_parts.append('</ul>')
                html_parts.append('</div>')
            
            # Реестр поручений — первый раздел отчёта
            action_registry = report_data.get('action_registry', '')
            if action_registry and action_registry.strip():
                registry_text = re.sub(r'(?m)^(#{1,5})\s', r'#\1 ', action_registry.strip())
                html_parts.append('<div class="section">')
                html_parts.append(f'<div class="content registry">{self._format_text_to_html(registry_text)}</div>')
                html_parts.append('</div>')

            # Протокол совещания
            protocol = report_data.get('protocol', '')
            if protocol and protocol.strip():
                html_parts.append('<div class="section">')
                html_parts.append(f'<div class="content protocol">{self._format_text_to_html(protocol)}</div>')
                html_parts.append('</div>')

            # Резюме
            summary = report_data.get('summary', '')
            if summary and summary.strip():
                html_parts.append('<div class="section">')
                html_parts.append('<h2>Резюме встречи</h2>')
                html_parts.append(f'<div class="content">{self._format_text_to_html(summary)}</div>')
                html_parts.append('</div>')
            
            # Открытые вопросы
            questions = report_data.get('questions', '')
            if questions and questions.strip():
                html_parts.append('<div class="section">')
                html_parts.append('<h2>Открытые вопросы</h2>')
                html_parts.append(f'<div class="content">{self._format_text_to_html(questions)}</div>')
                html_parts.append('</div>')

            participants_analysis = report_data.get('participants_analysis', '')
            if participants_analysis and participants_analysis.strip():
                html_parts.append('<div class="section">')
                html_parts.append('<h2>Анализ участников</h2>')
                html_parts.append(f'<div class="content">{self._format_text_to_html(participants_analysis)}</div>')
                html_parts.append('</div>')
            
            # Инсайты
            if insights:
                html_parts.append('<div class="section">')
                html_parts.append('<h2>Анализ и инсайты</h2>')
                
                # Динамика разговора
                dynamics = insights.get('conversation_dynamics', {})
                if dynamics:
                    html_parts.append('<h3>Динамика разговора</h3>')
                    html_parts.append('<div class="insights-grid">')
                    
                    balance = html.escape(dynamics.get('conversation_balance', ''))
                    interactivity = html.escape(dynamics.get('interactivity_level', ''))
                    switches = dynamics.get('total_speaker_switches', 0)
                    
                    html_parts.append(f'<div class="insight-item"><strong>Баланс разговора:</strong> {balance}</div>')
                    html_parts.append(f'<div class="insight-item"><strong>Интерактивность:</strong> {interactivity}</div>')
                    html_parts.append(f'<div class="insight-item"><strong>Переключения спикеров:</strong> {switches}</div>')
                    
                    html_parts.append('</div>')
                
                # Рекомендации
                recommendations = insights.get('recommendations', [])
                if recommendations:
                    html_parts.append('<h3>Рекомендации</h3>')
                    html_parts.append('<ul class="recommendations">')
                    
                    for rec in recommendations:
                        html_parts.append(f'<li>{html.escape(rec)}</li>')
                    
                    html_parts.append('</ul>')
                
                html_parts.append('</div>')
            
            # Полная транскрипция
            if transcript:
                html_parts.append('<div class="section">')
                html_parts.append('<h2>Полная транскрипция</h2>')
                html_parts.append('<div class="transcript">')
                
                for segment in transcript:
                    speaker = html.escape(segment.get('speaker', 'UNKNOWN'))
                    text = html.escape(segment.get('text', ''))
                    time_formatted = html.escape(segment.get('start_time_formatted', ''))
                    
                    html_parts.append('<div class="transcript-segment">')
                    html_parts.append(f'<div class="transcript-header">')
                    html_parts.append(f'<span class="speaker">{speaker}</span>')
                    html_parts.append(f'<span class="timestamp">{time_formatted}</span>')
                    html_parts.append('</div>')
                    html_parts.append(f'<div class="transcript-text">{text}</div>')
                    html_parts.append('</div>')
                
                html_parts.append('</div>')
                html_parts.append('</div>')
            
            # Закрытие HTML
            html_parts.append('</div>')  # container
            html_parts.append('</body>')
            html_parts.append('</html>')
            
            return '\n'.join(html_parts)
            
        except Exception as e:
            logger.error(f"❌ Ошибка построения HTML содержимого: {e}")
            return f'<html><body><h1>Ошибка генерации отчета</h1><p>Произошла ошибка: {html.escape(str(e))}</p></body></html>'
    
    def _format_text_to_html(self, text: str) -> str:
        """Форматирует текст для HTML с поддержкой Markdown-подобного синтаксиса."""
        try:
            blocks = self._parse_markdown_blocks(text)
            html_parts = []

            for block in blocks:
                if block['type'] == 'heading':
                    tag_level = min(6, block.get('level', 1) + 2)
                    html_parts.append(
                        f"<h{tag_level}>{self._format_inline_markdown_to_html(block['text'])}</h{tag_level}>"
                    )
                elif block['type'] == 'hr':
                    html_parts.append('<hr>')
                elif block['type'] == 'table':
                    html_parts.append('<table>')
                    html_parts.append('<thead><tr>' + ''.join(
                        f'<th>{self._format_inline_markdown_to_html(h)}</th>' for h in block['headers']
                    ) + '</tr></thead>')
                    html_parts.append('<tbody>')
                    for row in block['rows']:
                        html_parts.append('<tr>' + ''.join(
                            f'<td>{self._format_inline_markdown_to_html(c)}</td>' for c in row
                        ) + '</tr>')
                    html_parts.append('</tbody></table>')
                elif block['type'] == 'paragraph':
                    content = self._format_inline_markdown_to_html(block['text']).replace('\n', '<br>')
                    html_parts.append(f'<p>{content}</p>')
                elif block['type'] == 'ul':
                    html_parts.append('<ul>')
                    for item in block['items']:
                        html_parts.append(f"<li>{self._format_inline_markdown_to_html(item)}</li>")
                    html_parts.append('</ul>')
                elif block['type'] == 'ol':
                    html_parts.append('<ol>')
                    for item in block['items']:
                        html_parts.append(f"<li>{self._format_inline_markdown_to_html(item)}</li>")
                    html_parts.append('</ol>')

            return '\n'.join(html_parts)
            
        except Exception as e:
            logger.error(f"❌ Ошибка форматирования текста в HTML: {e}")
            return html.escape(text).replace('\n', '<br>')

    def _parse_markdown_blocks(self, text: str) -> List[Dict[str, Any]]:
        """Разбирает простой markdown-текст на блоки."""
        normalized = str(text).replace('\r\n', '\n').replace('\r', '\n')
        blocks: List[Dict[str, Any]] = []
        paragraph_lines: List[str] = []
        list_type: Optional[str] = None
        list_items: List[str] = []

        def flush_paragraph():
            nonlocal paragraph_lines
            if paragraph_lines:
                blocks.append({'type': 'paragraph', 'text': '\n'.join(paragraph_lines).strip()})
                paragraph_lines = []

        def flush_list():
            nonlocal list_type, list_items
            if list_type and list_items:
                blocks.append({'type': list_type, 'items': list_items[:]})
            list_type = None
            list_items = []

        def parse_table(start_idx: int) -> Optional[Dict[str, Any]]:
            """Пытается распарсить markdown-таблицу начиная со строки start_idx."""
            lines = normalized.split('\n')
            if start_idx + 1 >= len(lines):
                return None
            header_line = lines[start_idx].strip()
            sep_line = lines[start_idx + 1].strip()
            if not (header_line.startswith('|') and re.match(r'^\|[\s\-:|]+\|?\s*$', sep_line)):
                return None
            def cells(row: str) -> List[str]:
                return [c.strip() for c in row.strip().strip('|').split('|')]
            headers = cells(header_line)
            rows: List[List[str]] = []
            idx = start_idx + 2
            while idx < len(lines):
                row_line = lines[idx].strip()
                if row_line.startswith('|') and '|' in row_line[1:]:
                    rows.append(cells(row_line))
                    idx += 1
                else:
                    break
            return {'type': 'table', 'headers': headers, 'rows': rows, 'end_idx': idx}

        raw_lines = normalized.split('\n')
        i = 0
        while i < len(raw_lines):
            raw_line = raw_lines[i]
            line = raw_line.strip()

            if not line:
                flush_paragraph()
                flush_list()
                i += 1
                continue

            # Горизонтальный разделитель
            if re.match(r'^(-{3,}|\*{3,}|_{3,})$', line):
                flush_paragraph()
                flush_list()
                blocks.append({'type': 'hr'})
                i += 1
                continue

            # Markdown-таблица: строка |...| + строка-разделитель |---|
            if line.startswith('|'):
                flush_paragraph()
                flush_list()
                table = parse_table(i)
                if table:
                    blocks.append({'type': 'table', 'headers': table['headers'], 'rows': table['rows']})
                    i = table['end_idx']
                    continue
                # Не таблица — одиночная строка с пайпами, идёт в абзац
                paragraph_lines.append(line)
                i += 1
                continue

            heading_match = re.match(r'^(#{1,6})\s+(.*)$', line)
            if heading_match:
                flush_paragraph()
                flush_list()
                blocks.append({
                    'type': 'heading',
                    'level': len(heading_match.group(1)),
                    'text': heading_match.group(2).strip(),
                })
                i += 1
                continue

            unordered_match = re.match(r'^[-*]\s+(.*)$', line)
            if unordered_match:
                flush_paragraph()
                if list_type not in (None, 'ul'):
                    flush_list()
                list_type = 'ul'
                list_items.append(unordered_match.group(1).strip())
                i += 1
                continue

            ordered_match = re.match(r'^\d+\.\s+(.*)$', line)
            if ordered_match:
                flush_paragraph()
                if list_type not in (None, 'ol'):
                    flush_list()
                list_type = 'ol'
                list_items.append(ordered_match.group(1).strip())
                i += 1
                continue

            flush_list()
            paragraph_lines.append(line)
            i += 1

        flush_paragraph()
        flush_list()

        return blocks

    def _format_inline_markdown_to_html(self, text: str) -> str:
        """Рендерит inline markdown в HTML."""
        formatted = html.escape(str(text))
        links = []
        def _stash(m):
            links.append((m.group(1), m.group(2)))
            return f"\x00LINK{len(links) - 1}\x00"
        formatted = re.sub(r'\[([^\]]+)\]\(([^)\s]+)\)', _stash, formatted)
        formatted = re.sub(r'_{3,}', lambda m: '&#95;' * len(m.group(0)), formatted)
        formatted = re.sub(r'\*\*(.+?)\*\*', r'<strong>\1</strong>', formatted)
        formatted = re.sub(r'__(.+?)__', r'<strong>\1</strong>', formatted)
        formatted = re.sub(r'(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)', r'<em>\1</em>', formatted)
        formatted = re.sub(r'(?<!_)_(?!\s)(.+?)(?<!\s)_(?!_)', r'<em>\1</em>', formatted)
        formatted = re.sub(r'`([^`]+)`', r'<code>\1</code>', formatted)
        for i, (lt, lu) in enumerate(links):
            formatted = formatted.replace(f"\x00LINK{i}\x00", f'<a href="{lu}">{lt}</a>')
        return formatted
    
    def _get_html_styles(self) -> str:
        """Возвращает CSS стили для HTML отчета."""
        return """
        body {
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
            line-height: 1.6;
            color: #333;
            margin: 0;
            padding: 20px;
            background-color: #f5f5f5;
        }
        
        .container {
            max-width: 1200px;
            margin: 0 auto;
            background: white;
            padding: 30px;
            border-radius: 10px;
            box-shadow: 0 2px 10px rgba(0,0,0,0.1);
        }
        
        h1 {
            color: #2c3e50;
            border-bottom: 3px solid #3498db;
            padding-bottom: 10px;
            margin-bottom: 30px;
        }
        
        h2 {
            color: #34495e;
            margin-top: 30px;
            margin-bottom: 15px;
            border-left: 4px solid #3498db;
            padding-left: 15px;
        }
        
        h3 {
            color: #2c3e50;
            margin-top: 20px;
            margin-bottom: 10px;
        }

        h4, h5, h6 {
            color: #34495e;
            margin-top: 16px;
            margin-bottom: 8px;
        }
        
        .metadata {
            background: #ecf0f1;
            padding: 20px;
            border-radius: 8px;
            margin-bottom: 30px;
        }
        
        .info-grid {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(250px, 1fr));
            gap: 15px;
            margin-top: 15px;
        }
        
        .info-item {
            background: white;
            padding: 15px;
            border-radius: 5px;
            border-left: 3px solid #3498db;
        }
        
        .participants {
            margin-bottom: 30px;
        }
        
        .participants-grid {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(300px, 1fr));
            gap: 20px;
            margin-top: 15px;
        }
        
        .participant-card {
            background: #f8f9fa;
            padding: 20px;
            border-radius: 8px;
            border: 1px solid #dee2e6;
        }
        
        .participant-card h3 {
            margin-top: 0;
            color: #2c3e50;
        }
        
        .participation-bar {
            background: #e9ecef;
            height: 10px;
            border-radius: 5px;
            margin: 10px 0;
            overflow: hidden;
        }
        
        .participation-fill {
            background: linear-gradient(90deg, #3498db, #2ecc71);
            height: 100%;
            transition: width 0.3s ease;
        }
        
        .participant-stats {
            display: flex;
            justify-content: space-between;
            font-size: 0.9em;
            color: #6c757d;
        }
        
        .section {
            margin-bottom: 30px;
            padding: 20px;
            background: #fafafa;
            border-radius: 8px;
            border: 1px solid #e9ecef;
        }
        
        .content {
            margin-top: 15px;
        }

        .content p {
            margin: 0 0 12px 0;
        }

        .content ul,
        .content ol {
            margin: 8px 0 14px 20px;
            padding-left: 18px;
        }

        .content li {
            margin-bottom: 6px;
        }

        .content table {
            border-collapse: collapse;
            width: 100%;
            margin: 12px 0;
            font-size: 0.95em;
        }

        .content th,
        .content td {
            border: 1px solid #cbd5e1;
            padding: 8px 10px;
            text-align: left;
            vertical-align: top;
        }

        .content th {
            background: #3498db;
            color: white;
            font-weight: 600;
        }

        .content tr:nth-child(even) td {
            background: #f8fafc;
        }

        .content hr {
            border: none;
            border-top: 1px solid #cbd5e1;
            margin: 16px 0;
        }

        .content code {
            background: #eef2f7;
            border-radius: 4px;
            padding: 1px 5px;
            font-family: Consolas, Monaco, monospace;
            font-size: 0.95em;
        }
        
        .keywords {
            display: flex;
            flex-wrap: wrap;
            gap: 10px;
            margin-top: 15px;
        }
        
        .keyword {
            background: #3498db;
            color: white;
            padding: 5px 12px;
            border-radius: 20px;
            font-size: 0.9em;
            font-weight: 500;
        }
        
        .insights-grid {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(250px, 1fr));
            gap: 15px;
            margin-top: 15px;
        }
        
        .insight-item {
            background: white;
            padding: 15px;
            border-radius: 5px;
            border-left: 3px solid #e74c3c;
        }
        
        .recommendations {
            background: #d4edda;
            border: 1px solid #c3e6cb;
            border-radius: 5px;
            padding: 15px;
            margin-top: 15px;
        }
        
        .recommendations li {
            margin-bottom: 8px;
        }
        
        .transcript {
            max-height: 600px;
            overflow-y: auto;
            border: 1px solid #dee2e6;
            border-radius: 5px;
            margin-top: 15px;
        }
        
        .transcript-segment {
            padding: 15px;
            border-bottom: 1px solid #e9ecef;
        }
        
        .transcript-segment:last-child {
            border-bottom: none;
        }
        
        .transcript-header {
            display: flex;
            justify-content: space-between;
            align-items: center;
            margin-bottom: 8px;
        }
        
        .speaker {
            font-weight: bold;
            color: #2c3e50;
            background: #ecf0f1;
            padding: 3px 8px;
            border-radius: 3px;
        }
        
        .timestamp {
            font-size: 0.8em;
            color: #6c757d;
            font-family: monospace;
        }
        
        .transcript-text {
            color: #495057;
            line-height: 1.5;
        }
        
        @media (max-width: 768px) {
            .container {
                padding: 15px;
            }
            
            .info-grid,
            .participants-grid,
            .insights-grid {
                grid-template-columns: 1fr;
            }
            
            .participant-stats {
                flex-direction: column;
                gap: 5px;
            }
        }
        """
    
    def _generate_pdf_report(self, report_data: Dict[str, Any], html_path: Optional[str], base_filename: str) -> Optional[str]:
        """Генерирует PDF отчет из HTML или через fallback reportlab."""
        try:
            if html_path and Path(html_path).exists():
                try:
                    from weasyprint import HTML

                    output_path = self.output_dir / f"{base_filename}_report.pdf"
                    HTML(filename=html_path).write_pdf(str(output_path))

                    logger.info(f"✅ PDF отчет сохранен: {output_path}")
                    return str(output_path)

                except (ImportError, OSError) as exc:
                    logger.warning(f"⚠️ WeasyPrint недоступен ({exc}). Пробуем fallback через reportlab.")
            else:
                logger.warning("⚠️ HTML файл не найден для генерации PDF, используем fallback.")

            return self._generate_pdf_report_with_reportlab(report_data, base_filename)

        except Exception as e:
            logger.error(f"❌ Ошибка генерации PDF отчета: {e}")
            return None

    def _generate_pdf_report_with_reportlab(self, report_data: Dict[str, Any], base_filename: str) -> Optional[str]:
        """Генерирует PDF отчет через reportlab без внешних GTK-зависимостей."""
        try:
            from reportlab.lib import colors
            from reportlab.lib.enums import TA_CENTER
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer
        except ImportError:
            logger.warning("⚠️ ReportLab не установлен. PDF отчет не будет создан.")
            logger.info("💡 Для fallback PDF установите: pip install reportlab")
            return None

        try:
            regular_font_name = "Helvetica"
            bold_font_name = "Helvetica-Bold"
            regular_font_path = Path(r"C:\Windows\Fonts\arial.ttf")
            bold_font_path = Path(r"C:\Windows\Fonts\arialbd.ttf")

            if regular_font_path.exists() and bold_font_path.exists():
                regular_font_name = "VoiceVideoArial"
                bold_font_name = "VoiceVideoArialBold"
                if regular_font_name not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont(regular_font_name, str(regular_font_path)))
                if bold_font_name not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont(bold_font_name, str(bold_font_path)))
                pdfmetrics.registerFontFamily(
                    regular_font_name,
                    normal=regular_font_name,
                    bold=bold_font_name,
                    italic=regular_font_name,
                    boldItalic=bold_font_name,
                )

            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                "VoiceVideoTitle",
                parent=styles["Title"],
                fontName=bold_font_name,
                fontSize=18,
                leading=22,
                alignment=TA_CENTER,
                textColor=colors.HexColor("#1f2937"),
                spaceAfter=14,
            )
            heading_style = ParagraphStyle(
                "VoiceVideoHeading",
                parent=styles["Heading2"],
                fontName=bold_font_name,
                fontSize=13,
                leading=16,
                textColor=colors.HexColor("#111827"),
                spaceBefore=8,
                spaceAfter=6,
            )
            subheading_style = ParagraphStyle(
                "VoiceVideoSubheading",
                parent=heading_style,
                fontSize=11.5,
                leading=14,
                spaceBefore=5,
                spaceAfter=4,
            )
            body_style = ParagraphStyle(
                "VoiceVideoBody",
                parent=styles["BodyText"],
                fontName=regular_font_name,
                fontSize=9.5,
                leading=12,
                textColor=colors.HexColor("#374151"),
                spaceAfter=4,
            )
            meta_style = ParagraphStyle(
                "VoiceVideoMeta",
                parent=body_style,
                fontSize=10,
                leading=13,
                textColor=colors.HexColor("#4b5563"),
            )
            bullet_style = ParagraphStyle(
                "VoiceVideoBullet",
                parent=body_style,
                leftIndent=14,
                firstLineIndent=-8,
                spaceAfter=2,
            )

            output_path = self.output_dir / f"{base_filename}_report.pdf"
            doc = SimpleDocTemplate(
                str(output_path),
                pagesize=A4,
                leftMargin=16 * mm,
                rightMargin=16 * mm,
                topMargin=14 * mm,
                bottomMargin=14 * mm,
            )

            metadata = report_data.get("metadata", {})
            participants = report_data.get("participants", [])
            transcript = report_data.get("transcript", [])
            keywords = report_data.get("keywords", [])
            insights = report_data.get("insights", {})

            story = [
                Paragraph(self._pdf_text(metadata.get("title", "Анализ встречи")), title_style),
                Paragraph(
                    self._pdf_text(
                        f"Дата анализа: {metadata.get('generated_at', '')}\n"
                        f"Длительность: {metadata.get('total_duration', '')}\n"
                        f"Участников: {metadata.get('unique_speakers', 0)}\n"
                        f"Сегментов: {metadata.get('total_segments', 0)}"
                    ).replace("\n", "<br/>"),
                    meta_style,
                ),
                Spacer(1, 8),
            ]

            if participants:
                story.append(Paragraph("Участники", heading_style))
                for participant in participants:
                    story.append(
                        Paragraph(
                            f"• {self._pdf_text(participant.get('name', 'UNKNOWN'))}",
                            body_style,
                        )
                    )
                story.append(Spacer(1, 6))

            # Реестр поручений — первый раздел отчёта
            action_registry = report_data.get("action_registry", "")
            if action_registry and str(action_registry).strip():
                story.extend(
                    self._markdown_to_pdf_story(
                        str(action_registry),
                        body_style,
                        heading_style,
                        bullet_style,
                    )
                )
                story.append(Spacer(1, 6))

            protocol = report_data.get("protocol", "")
            if protocol and protocol.strip():
                story.extend(
                    self._markdown_to_pdf_story(
                        str(protocol),
                        body_style,
                        heading_style,
                        bullet_style,
                    )
                )
                story.append(Spacer(1, 6))

            for section_title, section_value in (
                ("Резюме встречи", report_data.get("summary", "")),
                ("Открытые вопросы", report_data.get("questions", "")),
                ("Анализ участников", report_data.get("participants_analysis", "")),
            ):
                if section_value:
                    story.append(Paragraph(section_title, heading_style))
                    story.extend(
                        self._markdown_to_pdf_story(
                            str(section_value),
                            body_style,
                            subheading_style,
                            bullet_style,
                        )
                    )
                    story.append(Spacer(1, 6))

            if insights:
                story.append(Paragraph("Инсайты", heading_style))
                story.extend(
                    self._markdown_to_pdf_story(
                        json.dumps(insights, ensure_ascii=False, indent=2),
                        body_style,
                        subheading_style,
                        bullet_style,
                    )
                )
                story.append(Spacer(1, 6))

            if transcript:
                story.append(PageBreak())
                story.append(Paragraph("Транскрипция", heading_style))
                for segment in transcript:
                    line = (
                        f"[{segment.get('start_time_formatted', '00:00')} - "
                        f"{segment.get('end_time_formatted', '00:00')}] "
                        f"<b>{self._pdf_text(segment.get('speaker', 'UNKNOWN'))}</b>: "
                        f"{self._pdf_text(segment.get('text', ''))}"
                    )
                    story.append(Paragraph(line, body_style))
                    story.append(Spacer(1, 3))

            doc.build(story)
            logger.info(f"✅ PDF отчет сохранен через reportlab: {output_path}")
            return str(output_path)

        except Exception as e:
            logger.error(f"❌ Ошибка fallback PDF через reportlab: {e}")
            return None

    def _pdf_paragraphs(self, text: str, style) -> List[Any]:
        """Разбивает длинный текст на Paragraph-блоки для PDF."""
        from reportlab.platypus import Paragraph

        normalized = str(text).replace("\r\n", "\n").replace("\r", "\n")
        blocks = [block.strip() for block in normalized.split("\n\n") if block.strip()]
        if not blocks:
            return [Paragraph(self._pdf_text(normalized), style)]
        return [Paragraph(self._pdf_text(block).replace("\n", "<br/>"), style) for block in blocks]

    def _markdown_to_pdf_story(self, text: str, body_style, heading_style, bullet_style) -> List[Any]:
        """Рендерит markdown-текст в список flowables для reportlab."""
        from reportlab.lib import colors
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import HRFlowable, Paragraph, Spacer, Table, TableStyle

        story: List[Any] = []
        for block in self._parse_markdown_blocks(text):
            if block['type'] == 'heading':
                story.append(Paragraph(self._format_inline_markdown_to_pdf(block['text']), heading_style))
                story.append(Spacer(1, 2))
            elif block['type'] == 'hr':
                story.append(HRFlowable(width='100%', thickness=0.6, color=colors.HexColor('#94a3b8')))
            elif block['type'] == 'table':
                cell_style = ParagraphStyle(
                    'TableCell',
                    parent=body_style,
                    fontSize=8.5,
                    leading=10.5,
                    spaceAfter=0,
                )
                header_style = ParagraphStyle(
                    'TableHeader',
                    parent=cell_style,
                    textColor=colors.white,
                )
                data = [[Paragraph(f"<b>{self._format_inline_markdown_to_pdf(h)}</b>", header_style)
                         for h in block['headers']]]
                for row in block['rows']:
                    data.append([Paragraph(self._format_inline_markdown_to_pdf(c), cell_style)
                                 for c in row])
                col_count = max(len(block['headers']), 1)
                avail_width = 178 * mm
                col_widths = [avail_width / col_count] * col_count
                # Узкая колонка "№" — если первая колонка короткая
                if col_count > 1:
                    first_col_max = max(
                        [len(str(block['headers'][0]))] +
                        [len(str(r[0])) for r in block['rows'] if r]
                    )
                    if first_col_max <= 4:
                        col_widths[0] = 10 * mm
                        rest = (avail_width - 10 * mm) / (col_count - 1)
                        col_widths[1:] = [rest] * (col_count - 1)
                table = Table(data, colWidths=col_widths, repeatRows=1)
                table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3498db')),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#94a3b8')),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('LEFTPADDING', (0, 0), (-1, -1), 4),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                    ('TOPPADDING', (0, 0), (-1, -1), 3),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f1f5f9')]),
                ]))
                story.append(table)
                story.append(Spacer(1, 6))
            elif block['type'] == 'paragraph':
                story.append(
                    Paragraph(
                        self._format_inline_markdown_to_pdf(block['text']).replace('\n', '<br/>'),
                        body_style,
                    )
                )
            elif block['type'] == 'ul':
                for item in block['items']:
                    story.append(Paragraph(f"• {self._format_inline_markdown_to_pdf(item)}", bullet_style))
            elif block['type'] == 'ol':
                for index, item in enumerate(block['items'], start=1):
                    story.append(Paragraph(f"{index}. {self._format_inline_markdown_to_pdf(item)}", bullet_style))
            story.append(Spacer(1, 2))

        return story or [Paragraph(self._pdf_text(text), body_style)]

    def _format_inline_markdown_to_pdf(self, text: str) -> str:
        """Рендерит inline markdown в теги, поддерживаемые reportlab Paragraph."""
        formatted = self._pdf_text(text)
        links = []
        def _stash(m):
            links.append((m.group(1), m.group(2)))
            return f"\x00LINK{len(links) - 1}\x00"
        formatted = re.sub(r'\[([^\]]+)\]\(([^)\s]+)\)', _stash, formatted)
        # Длинные подчёркивания (линии подписи ____) защищаем от разбора как курсив/болд
        formatted = re.sub(r'_{3,}', lambda m: '&#95;' * len(m.group(0)), formatted)
        formatted = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', formatted)
        formatted = re.sub(r'__(.+?)__', r'<b>\1</b>', formatted)
        formatted = re.sub(r'(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)', r'<i>\1</i>', formatted)
        formatted = re.sub(r'(?<!_)_(?!\s)(.+?)(?<!\s)_(?!_)', r'<i>\1</i>', formatted)
        formatted = re.sub(r'`([^`]+)`', r'<font face="Courier">\1</font>', formatted)
        for i, (lt, lu) in enumerate(links):
            formatted = formatted.replace(
                f"\x00LINK{i}\x00",
                f'<link href="{lu}" color="#1d4ed8"><u>{lt}</u></link>'
            )
        # Остаточные непарные маркеры — убираем, чтобы в PDF не было «**» и «__»
        formatted = formatted.replace('**', '').replace('__', '')
        return formatted

    def _pdf_text(self, text: Any) -> str:
        """Экранирует текст для reportlab Paragraph."""
        return xml_escape(str(text), {"\"": "&quot;"})

    def _generate_docx_reports(self, report_data: Dict[str, Any], base_filename: str) -> Dict[str, str]:
        """Генерирует DOCX: полный отчет и отдельный файл протокола."""
        files: Dict[str, str] = {}
        try:
            import docx  # noqa: F401
        except ImportError:
            logger.warning("⚠️ python-docx не установлен. DOCX отчеты не будут созданы.")
            return files

        try:
            metadata = report_data.get("metadata", {})
            participants = report_data.get("participants", [])

            # Полный отчет
            doc = self._new_docx_document()
            doc.add_heading(metadata.get("title", "Анализ встречи"), level=0)

            doc.add_heading("Информация о встрече", level=1)
            for label, value in (
                ("Дата анализа", metadata.get("generated_at", "")),
                ("Длительность", metadata.get("total_duration", "")),
                ("Количество участников", metadata.get("unique_speakers", 0)),
                ("Всего сегментов", metadata.get("total_segments", 0)),
            ):
                p = doc.add_paragraph()
                p.add_run(f"{label}: ").bold = True
                p.add_run(str(value))

            if participants:
                doc.add_heading("Участники", level=1)
                for participant in participants:
                    doc.add_paragraph(
                        participant.get("name", "Неизвестный"), style="List Bullet"
                    )

            # Реестр поручений — первый раздел отчёта
            action_registry = report_data.get("action_registry", "")
            if action_registry and str(action_registry).strip():
                self._append_markdown_to_docx(doc, str(action_registry))

            protocol = report_data.get("protocol", "")
            if protocol and protocol.strip():
                self._append_markdown_to_docx(doc, protocol)

            for section_title, section_value in (
                ("Резюме встречи", report_data.get("summary", "")),
                ("Открытые вопросы", report_data.get("questions", "")),
                ("Анализ участников", report_data.get("participants_analysis", "")),
            ):
                if section_value and str(section_value).strip():
                    doc.add_heading(section_title, level=1)
                    self._append_markdown_to_docx(doc, str(section_value))

            transcript = report_data.get("transcript", [])
            if transcript:
                doc.add_page_break()
                doc.add_heading("Полная транскрипция", level=1)
                for segment in transcript:
                    p = doc.add_paragraph()
                    ts = segment.get("start_time_formatted", "")
                    p.add_run(f"[{ts}] ").bold = True
                    p.add_run(f"{segment.get('speaker', 'UNKNOWN')}: ").bold = True
                    p.add_run(segment.get("text", ""))

            report_path = self.output_dir / f"{base_filename}_report.docx"
            doc.save(str(report_path))
            files["docx"] = str(report_path)
            logger.info(f"✅ DOCX отчет сохранен: {report_path}")

            # Отдельный файл протокола (реестр поручений первым — подписной документ)
            if protocol and protocol.strip():
                protocol_doc = self._new_docx_document()
                if action_registry and str(action_registry).strip():
                    self._append_markdown_to_docx(protocol_doc, str(action_registry))
                self._append_markdown_to_docx(protocol_doc, protocol)
                protocol_path = self.output_dir / f"{base_filename}_protocol.docx"
                protocol_doc.save(str(protocol_path))
                files["protocol_docx"] = str(protocol_path)
                logger.info(f"✅ DOCX протокол сохранен: {protocol_path}")

            return files

        except Exception as e:
            logger.error(f"❌ Ошибка генерации DOCX: {e}")
            return files

    def _new_docx_document(self):
        """Создает Document с кириллическим базовым шрифтом."""
        import docx
        from docx.shared import Pt

        document = docx.Document()
        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(10)
        return document

    def _append_markdown_to_docx(self, document, text: str) -> None:
        """Рендерит markdown-блоки в документ Word (заголовки, списки, таблицы)."""
        for block in self._parse_markdown_blocks(text):
            if block["type"] == "heading":
                document.add_heading(
                    self._strip_markdown(block["text"]),
                    level=min(9, block.get("level", 1) + 1),
                )
            elif block["type"] == "paragraph":
                for line in block["text"].split("\n"):
                    paragraph = document.add_paragraph()
                    self._add_md_runs_to_docx(paragraph, line)
            elif block["type"] == "ul":
                for item in block["items"]:
                    paragraph = document.add_paragraph(style="List Bullet")
                    self._add_md_runs_to_docx(paragraph, item)
            elif block["type"] == "ol":
                for item in block["items"]:
                    paragraph = document.add_paragraph(style="List Number")
                    self._add_md_runs_to_docx(paragraph, item)
            elif block["type"] == "table":
                headers = block["headers"]
                table = document.add_table(rows=len(block["rows"]) + 1, cols=len(headers))
                table.style = "Table Grid"
                for col_idx, header in enumerate(headers):
                    cell = table.rows[0].cells[col_idx]
                    cell.paragraphs[0].add_run(self._strip_markdown(header)).bold = True
                for row_idx, row in enumerate(block["rows"], start=1):
                    for col_idx, cell_text in enumerate(row[: len(headers)]):
                        cell = table.rows[row_idx].cells[col_idx]
                        self._add_md_runs_to_docx(cell.paragraphs[0], cell_text)
            # hr — в Word не нужен, пропускаем

    def _add_hyperlink_to_docx(self, paragraph, text: str, url: str) -> None:
        """Вставляет кликабельную гиперссылку в параграф Word."""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.opc.constants import RELATIONSHIP_TYPE
        r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "1D4ED8")
        rpr.append(color)
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        rpr.append(underline)
        run.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    def _add_md_runs_to_docx(self, paragraph, text: str) -> None:
        """Добавляет runs в параграф Word с разбором **bold** и [текст](ссылка)."""
        parts = re.split(r"(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\))", str(text))
        for part in parts:
            if not part:
                continue
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            if link:
                self._add_hyperlink_to_docx(paragraph, link.group(1), link.group(2))
            elif part.startswith("**") and part.endswith("**") and len(part) > 4:
                run = paragraph.add_run(self._strip_markdown(part[2:-2]))
                run.bold = True
            else:
                paragraph.add_run(self._strip_markdown(part))

    def _strip_markdown(self, text: str) -> str:
        """Убирает inline-разметку markdown из строки."""
        cleaned = re.sub(r"\*\*(.+?)\*\*", r"\1", str(text))
        cleaned = re.sub(r"__(.+?)__", r"\1", cleaned)
        cleaned = re.sub(r"`([^`]+)`", r"\1", cleaned)
        cleaned = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", cleaned)
        return cleaned.replace("**", "").replace("__", "")

    def _generate_transcript_file(self, segments_with_speakers: List[Dict], base_filename: str,
                                  word_segments: Optional[List[Dict]] = None) -> Optional[str]:
        """Генерирует отдельный файл транскрипции в JSON формате."""
        try:
            output_path = self.output_dir / f"{base_filename}_transcript.json"

            transcript_data = {
                'metadata': {
                    'generated_at': datetime.now().isoformat(),
                    'total_segments': len(segments_with_speakers),
                    'format': 'detailed_transcript'
                },
                'segments': self._format_transcript(segments_with_speakers)
            }
            if word_segments:
                transcript_data['word_segments'] = word_segments
            
            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(transcript_data, f, ensure_ascii=False, indent=2)
            
            logger.info(f"✅ Файл транскрипции сохранен: {output_path}")
            return str(output_path)
            
        except Exception as e:
            logger.error(f"❌ Ошибка генерации файла транскрипции: {e}")
            return None
